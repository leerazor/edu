"""Validate both separate story revisions after actual PowerPoint rendering.

Run from the repository root after build_story_pair.py and
render_story_pair_slides.sh. Reports and contact sheets are scoped to these
versions. External reference attribution and asset hashes are checked too.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from pathlib import Path

import fitz
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation

from presentation_common import FONT
from build_story_pair import check_source
from production_common import (
    compact,
    cue_text,
    markdown_scene_blocks,
    measure,
    name,
    prompter_scene_blocks,
    prompter_turns,
    scene_content_hash,
    timecode,
)

ROOT = Path(__file__).resolve().parents[1]
PAGE_IDS = [f"P{i:02d}" for i in range(1, 11)] + [f"A{i:02d}" for i in range(1, 4)]
ALLOWED_FONTS = {FONT, "Malgun Gothic", "맑은 고딕"}
EXPECTED_CANDIDATES = {
    "seminar": "AI 4 AI 세미나", "tech-x": "Tech X",
    "job-fair": "Job Fair", "dcis": "DCIS",
}
EXPECTED_MATCHES = [
    ("M01", ["seminar", "tech-x"], "tech-x"),
    ("M02", ["job-fair", "dcis"], "dcis"),
    ("M03", ["tech-x", "dcis"], "tech-x"),
]


class Validation:
    def __init__(self, package):
        self.package = package
        self.out = ROOT / "deliverables" / package
        self.docs = ROOT / "docs" / package
        self.passed = []
        self.errors = []
        self.metrics = {}

    def check(self, condition, message):
        (self.passed if condition else self.errors).append(message)

    def section(self, label, action):
        try:
            return action()
        except Exception as error:
            self.errors.append(f"{label}: {type(error).__name__}: {error}")
            return None


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_facts(film, validation):
    check = validation.check
    scenes = film["scenes"]
    spoken = " ".join(t["text"] for s in scenes for t in s["turns"])
    introduction = " ".join(t["text"] for s in scenes[:3] for t in s["turns"])
    intro_material = introduction + json.dumps([s["visual"] for s in scenes[:3]], ensure_ascii=False)
    check(all(term in introduction for term in ["DS부문", "모든", "인프라", "AI"]),
          "Opening states the DS infrastructure role and AI direction")
    check(all(term in intro_material for term in ["7만", "32", "100만"]),
          "Opening uses the newly supplied team scale figures")
    check(not re.search(r"160|백육십", intro_material), "Opening omits team headcount emphasis")
    check(any(term in introduction for term in ["교류", "경험을 나누", "경험을 나눕"]) and "성장" in introduction,
          "Opening connects exchange with growth")
    tech_x_material = spoken + json.dumps([s["visual"] for s in scenes], ensure_ascii=False)
    check("3월" in tech_x_material and "4월" in tech_x_material,
          "Tech X recruitment and project start are distinct")
    check("Advanced Infrastructure for Artificial Intelligence" in
          json.dumps(film, ensure_ascii=False), "Official AI 4 AI expansion retained")
    dates = {"outsystems": "2026-10-12", "submission": "2026-10-14",
             "tech_x_final": "2026-10-16", "filming": "2026-10-19"}
    check(film["schedule"]["as_of"] == "2026-10-11", "Planning status dated October 11")
    for key, date in dates.items():
        check(film["schedule"][key] == {"date": date, "status": "scheduled"},
              f"{key} date and scheduled status")
    iap_scenes = [s for s in scenes if re.search(
        r"(?<![A-Za-z])IAP(?![A-Za-z])|Infra Agent Platform",
        " ".join(t["text"] for t in s["turns"]))]
    expected_iap = "S06" if validation.package == "story-v02" else "S05"
    check([s["id"] for s in iap_scenes] == [expected_iap], "IAP is one brief Tech X example")
    check(film.get("iap_scene") == expected_iap, "IAP scene metadata matches speech")
    check(all(s["duration"] <= 20 for s in iap_scenes), "IAP example has at most 20 seconds")
    gpu = scenes[4] if validation.package == "story-v02" else scenes[3]
    check({"IMG16", "IMG17"}.issubset(gpu["assets"]), "Personal GPU experience references both user photo slots")
    check({"IMG16", "IMG17"}.issubset({slot["id"] for slot in gpu.get("asset_slots", [])}),
          "Both GPU photos are editable material slots")
    validation.metrics["iap_seconds"] = sum(s["duration"] for s in iap_scenes)


def check_timing_and_documents(film, validation):
    check = validation.check
    check_source(film)
    check(True, "Source scene continuity, speakers, cues, and allocation")
    scenes = film["scenes"]
    rows = [measure(s, film) for s in scenes]
    timing = json.loads((validation.out / "timing-report.json").read_text(encoding="utf-8"))
    check(timing["scenes"] == rows, "Timing report matches current pronunciation estimates")
    check(timing["version"] == film["version"] and timing["date"] == film["date"],
          "Timing report version and date")
    check(timing["main_allocated_seconds"] == 275 and timing["fixed_seconds"] == 25,
          "Main 275 seconds and fixed footage 25 seconds")
    check(timing["reading_rates_syllables_per_second"] == film["reading_rates"],
          "Reading rate assumptions match source")
    check(timing["total_spoken_syllables"] == sum(row["spoken_syllables"] for row in rows),
          "Total spoken syllables")
    check(timing["iap_allocated_seconds"] == validation.metrics["iap_seconds"],
          "Short IAP allocation agrees with timing report")
    for rate in film["reading_rates"]:
        total = round(sum(row["estimates_seconds"][rate] for row in rows), 2)
        check(timing["main_estimates_seconds"][rate] == total and
              timing["whole_film_estimates_seconds"][rate] == round(total + 25, 2),
              f"{rate} main and whole estimates")
    check(timing["whole_film_estimates_seconds"]["slow"] < 600, "Usable take estimate below ten minutes")
    script = (validation.docs / "02-script.md").read_text(encoding="utf-8")
    board = (validation.docs / "01-storyboard.md").read_text(encoding="utf-8")
    prompt = (validation.out / "teleprompter.txt").read_text(encoding="utf-8")
    script_blocks, board_blocks = markdown_scene_blocks(script), markdown_scene_blocks(board)
    prompt_blocks = prompter_scene_blocks(prompt)
    expected_ids = {s["id"] for s in scenes}
    check(set(script_blocks) == set(board_blocks) == set(prompt_blocks) == expected_ids,
          "Script, storyboard, and prompter contain the ten current scenes")
    facts = (ROOT / "docs/03-facts-and-assets.md").read_text(encoding="utf-8")
    registered = set(re.findall(r"(?m)^\|\s*(IMG\d{2})\s*\|", facts))
    for doc in validation.docs.glob("*.md"):
        value = doc.read_text(encoding="utf-8")
        check("이 문서에서 해야 할 일" in value, f"{doc.name}: reader action")
        check(set(re.findall(r"\bIMG\d{2}\b", value)).issubset(registered),
              f"{doc.name}: registered internal asset IDs")
    with (validation.out / "rehearsal.csv").open(encoding="utf-8-sig", newline="") as handle:
        rehearsal = list(csv.DictReader(handle))
    check([r["scene"] for r in rehearsal] == [s["id"] for s in scenes], "Rehearsal row order")
    for scene, row, estimate in zip(scenes, rehearsal, rows, strict=True):
        sid, pid = scene["id"], scene["slide"]
        check(set(scene["assets"]).issubset(registered), f"{sid}: registered asset references")
        check(estimate["estimates_seconds"]["slow"] <= scene["duration"], f"{sid}: slow estimate within allocation")
        check(set(re.findall(r"P\d{2}", scene["cue"]["effect"])) == {pid}, f"{sid}: main cue points to its slide")
        check(scene["camera"] in board_blocks[sid] and scene["hold"] in board_blocks[sid],
              f"{sid}: storyboard directions match source")
        check(scene["cue"]["anchor"] in board_blocks[sid], f"{sid}: storyboard cue")
        for turn in scene["turns"]:
            expected = f"{name(turn['speaker'], film)}: {cue_text(turn, scene, film)}"
            check(compact(expected) in compact(script_blocks[sid].replace("**", "")),
                  f"{sid}: script speech and cue")
        for slot in scene.get("asset_slots", []):
            labels = [slot["label"], slot["label"].replace("\n", " / ")]
            check(any(label in script_blocks[sid] for label in labels) and slot["detail"] in board_blocks[sid],
                  f"{sid}: material insertion instructions in documents")
        actual_slide, block = prompt_blocks[sid]
        expected_turns = [(name(t["speaker"], film), compact(cue_text(t, scene, film, phonetic=True)))
                          for t in scene["turns"]]
        check(actual_slide == pid and
              [(who, compact(value)) for who, value in prompter_turns(block)] == expected_turns,
              f"{sid}: prompter pronunciation and cue")
        check(row["content_hash"] == scene_content_hash(scene, film), f"{sid}: rehearsal content hash")
        check(row["PPT"] == pid and row["start"] == timecode(scene["start"]) and
              row["end"] == timecode(scene["start"] + scene["duration"]) and
              float(row["allocated_seconds"]) == scene["duration"], f"{sid}: rehearsal slide and times")
        check(int(row["syllables"]) == estimate["spoken_syllables"] and
              float(row["estimated_baseline_seconds"]) == estimate["estimates_seconds"]["baseline"],
              f"{sid}: rehearsal estimate")
        if row["measured_seconds"].strip():
            check(float(row["measured_seconds"]) > 0, f"{sid}: positive rehearsal measurement")
    if validation.package == "worldcup-v02":
        for label, value in [("script", script), ("storyboard", board)]:
            check(film["tournament"]["disclaimer"] in value, f"Tournament {label}: proposed selection disclosed")
            check(film["tournament"]["criterion"] in value, f"Tournament {label}: shared selection criterion")
    validation.metrics.update({"main_seconds": 275, "whole_seconds": 300,
        "spoken_syllables": timing["total_spoken_syllables"],
        "estimates": timing["whole_film_estimates_seconds"],
        "measured_scenes": sum(bool(r["measured_seconds"].strip()) for r in rehearsal),
        "asset_slots": sum(len(s.get("asset_slots", [])) for s in scenes)})


def check_deck_source(film, deck, validation):
    check = validation.check
    check(len(deck.slides) == 13, "Ten main slides and three appendices")
    for index, slide in enumerate(deck.slides):
        pid = PAGE_IDS[index]
        check(slide._element.get("show", "1") == ("1" if index < 10 else "0"), f"{pid}: visibility")
        notes = slide.notes_slide.notes_text_frame.text
        check(bool(notes), f"{pid}: notes")
        if index >= 10:
            continue
        scene = film["scenes"][index]
        check(scene["camera"] in notes and scene["hold"] in notes, f"{pid}: notes directions")
        for turn in scene["turns"]:
            check(cue_text(turn, scene, film) in notes, f"{pid}: notes speech and cue")
        for cue in scene.get("secondary_cues", []):
            check(all(cue[key] in notes for key in ["anchor", "label", "effect"]), f"{pid}: secondary cue")
        shapes = {shape.name: shape for shape in slide.shapes}
        for slot in scene.get("asset_slots", []):
            frame = shapes.get(f"{pid}-asset-frame-{slot['id']}")
            label = shapes.get(f"{pid}-asset-label-{slot['id']}")
            check(frame is not None, f"{pid}: editable frame {slot['id']}")
            check(label is not None and label.text == slot["label"], f"{pid}: material label {slot['id']}")
            check(all(slot[key] in notes for key in ["id", "label", "detail"]), f"{pid}: material notes {slot['id']}")
    for scene in film["scenes"]:
        pid = scene["slide"]
        slide = deck.slides[int(pid[1:]) - 1]
        visible = compact("\n".join(shape.text for shape in slide.shapes if shape.has_text_frame))
        visual = scene["visual"]
        for key in ["eyebrow", "headline", "body", "footer", "bottom", "left", "right", "criterion", "detail"]:
            if visual.get(key):
                check(compact(visual[key]) in visible, f"{pid}: source {key} appears in deck")
        for item in visual.get("items", []):
            for key in ["label", "value", "detail"]:
                if item.get(key):
                    check(compact(item[key]) in visible, f"{pid}: source item {key} appears in deck")


def contact_sheet(preview, ids, validation, choice=False):
    columns = 3
    rows = math.ceil(len(ids) / columns)
    canvas = Image.new("RGB", (1640, rows * 340 + 30), "#E8EBE5")
    draw = ImageDraw.Draw(canvas)
    for index, pid in enumerate(ids):
        x, y = 20 + (index % columns) * 540, 20 + (index // columns) * 340
        with Image.open(preview / f"slide-{index + 1:02}.png") as image:
            canvas.paste(image.convert("RGB").resize((520, 293), Image.Resampling.LANCZOS), (x, y))
        label = pid if choice else pid + (" / MAIN" if index < 10 else " / HIDDEN APPENDIX")
        draw.text((x, y + 303), label, fill="#202522")
    canvas.save(preview / "contact-sheet.jpg", quality=93)
    validation.check(True, "Contact sheet generated for visual review")


def check_render(deck_path, pdf_path, preview, ids, validation, choice=False):
    check = validation.check
    with zipfile.ZipFile(deck_path) as archive:
        check(archive.testzip() is None, f"{deck_path.name}: PPTX archive integrity")
        slide_xml = [p for p in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", p)]
        check(not any(b"<p:timing" in archive.read(p) for p in slide_xml),
              f"{deck_path.name}: static reference deck")
    deck = Presentation(deck_path)
    check(abs(deck.slide_width / deck.slide_height - 16 / 9) < .001, f"{deck_path.name}: 16:9 aspect ratio")
    check(len(deck.slides) == len(ids), f"{deck_path.name}: expected slide count")
    for pid, slide in zip(ids, deck.slides, strict=True):
        for shape in slide.shapes:
            check(shape.left >= 0 and shape.top >= 0 and
                  shape.left + shape.width <= deck.slide_width + 20 and
                  shape.top + shape.height <= deck.slide_height + 20, f"{pid}: shape bounds")
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            check(run.font.name in ALLOWED_FONTS and run.font.size is not None,
                                  f"{pid}: explicit font")
    check({p.name for p in preview.glob("slide-*.png")} ==
          {f"slide-{i:02}.png" for i in range(1, len(ids) + 1)}, f"{deck_path.name}: exact PNG export set")
    with fitz.open(pdf_path) as pdf:
        check(len(pdf) == len(ids), f"{deck_path.name}: PDF page count")
        for index, page in enumerate(pdf):
            pid = ids[index]
            extracted = compact(page.get_text())
            check("�" not in extracted, f"{pid}: PDF encoding")
            if not choice:
                check(pid in extracted, f"{pid}: PDF page ID")
            for shape in deck.slides[index].shapes:
                if shape.has_text_frame and shape.text.strip():
                    check(compact(shape.text) in extracted, f"{pid}: all visible text rendered ({shape.text[:20]})")
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        check(x0 >= -1 and y0 >= -1 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1,
                              f"{pid}: PDF text bounds")
            with Image.open(preview / f"slide-{index + 1:02}.png") as image:
                check(image.size == (1600, 900), f"{pid}: PNG dimensions")
                pix = page.get_pixmap(matrix=fitz.Matrix(.2, .2), alpha=False)
                ref = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                a = image.convert("L").resize((96, 54), Image.Resampling.LANCZOS)
                b = ref.convert("L").resize((96, 54), Image.Resampling.LANCZOS)
                check(ImageStat.Stat(ImageChops.difference(a, b)).mean[0] / 255 < .12,
                      f"{pid}: PNG and PDF agree")
    contact_sheet(preview, ids, validation, choice=choice)
    return deck


def check_tournament(film, validation):
    check = validation.check
    tournament = film["tournament"]
    check(tournament["actual_vote"] is False and tournament["status"] == "촬영용 추천안",
          "Tournament is a filming proposal rather than an actual vote")
    check(bool(tournament["criterion"].strip()), "One shared selection criterion")
    roster = {c["id"]: c["name"] for c in tournament["candidates"]}
    check(roster == EXPECTED_CANDIDATES and len(tournament["candidates"]) == 4, "Four candidates include DCIS")
    matches = tournament["matches"]
    check([(m["id"], m["candidates"], m["winner"]) for m in matches] == EXPECTED_MATCHES,
          "Bracket progression matches the proposed Tech X / DCIS final")
    for match in matches:
        if "criterion" in match:
            check(match["criterion"] == tournament["criterion"], f"{match['id']}: same selection criterion")
        reveal_scene = next(s for s in film["scenes"] if s["id"] == match["reveal_scene"])
        reveal_speech = " ".join(turn["text"] for turn in reveal_scene["turns"])
        check(reveal_speech.count(match["reveal_anchor"]) == 1,
              f"{match['id']}: proposed choice cue occurs once in its reveal scene")
        check(roster[match["winner"]] in match["reveal_anchor"],
              f"{match['id']}: reveal speech agrees with proposed winner")
    disclaimer = tournament["disclaimer"]
    check(all(word in disclaimer for word in ["촬영", "선택", "실제", "투표"]),
          "Proposal and actual vote distinction disclosed")
    deck_path = validation.out / "choice-states.pptx"
    choice_ids = [f"{match['id']}-{state}" for match in matches for state in ["before", "after"]]
    choice = check_render(deck_path, validation.out / "choice-states.pdf",
                          validation.out / "choice-preview", choice_ids, validation, choice=True)
    for index, slide in enumerate(choice.slides):
        match = matches[index // 2]
        state = "선택 전" if index % 2 == 0 else "선택 후"
        notes = slide.notes_slide.notes_text_frame.text
        visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        check(match["id"] in notes and compact(state) in compact(notes), f"{choice_ids[index]}: match and state notes")
        check(tournament["criterion"] in notes, f"{choice_ids[index]}: selection criterion notes")
        check(match["winner"] in notes or roster[match["winner"]] in notes, f"{choice_ids[index]}: winner notes")
        check(all(roster[candidate] in visible for candidate in match["candidates"]),
              f"{choice_ids[index]}: both candidate names visible")
        if index % 2:
            check(roster[match["winner"]] in visible, f"{choice_ids[index]}: winner visible after reveal")
    for index, match in enumerate(matches):
        before = choice.slides[index * 2]
        after = choice.slides[index * 2 + 1]
        before_shapes = b"".join(shape._element.xml.encode("utf-8") for shape in before.shapes)
        after_shapes = b"".join(shape._element.xml.encode("utf-8") for shape in after.shapes)
        check(before_shapes != after_shapes, f"{match['id']}: reveal changes visible slide objects")
    validation.metrics["selection_status"] = tournament["status"]
    validation.metrics["actual_vote"] = False


def check_reference_assets(validation):
    manifest = ROOT / "assets/references/manifest.json"
    for entry in json.loads(manifest.read_text(encoding="utf-8")):
        path = ROOT / entry["file"]
        validation.check(path.is_file() and sha256(path) == entry["sha256"], f"Attributed reference asset: {entry['file']}")


def validate(package):
    validation = Validation(package)
    source = ROOT / "content" / f"{package}-film.json"
    film = validation.section("Source JSON", lambda: json.loads(source.read_text(encoding="utf-8")))
    if film is not None:
        validation.section("Confirmed facts and scope", lambda: check_facts(film, validation))
        validation.section("Timing and document synchronization", lambda: check_timing_and_documents(film, validation))
        deck_path = validation.out / f"learning-fair-2026-{package}.pptx"
        deck = validation.section("PowerPoint render", lambda: check_render(deck_path,
            validation.out / f"learning-fair-2026-{package}.pdf", validation.out / "preview", PAGE_IDS, validation))
        if deck is not None:
            validation.section("Deck source synchronization", lambda: check_deck_source(film, deck, validation))
        if package == "worldcup-v02":
            validation.section("Tournament and reveal states", lambda: check_tournament(film, validation))
    validation.section("Attributed external assets", lambda: check_reference_assets(validation))
    artifacts = [source, validation.out / f"learning-fair-2026-{package}.pptx",
        validation.out / f"learning-fair-2026-{package}.pdf", validation.docs / "01-storyboard.md",
        validation.docs / "02-script.md", validation.out / "teleprompter.txt", validation.out / "timing-report.json"]
    if package == "worldcup-v02":
        artifacts += [validation.out / "choice-states.pptx", validation.out / "choice-states.pdf"]
    report = {"status": "FAIL" if validation.errors else "PASS", "passed_checks": len(validation.passed),
        "errors": validation.errors, "metrics": validation.metrics,
        "artifact_sha256": {str(p.relative_to(ROOT)): sha256(p) for p in artifacts if p.is_file()},
        "limits": ["Reading estimates are not presenter rehearsal measurements.",
            "Actual internal photos and company KV will be supplied separately.",
            "Contact sheets require visual review. Motion is provided as editing instructions.",
            "Tournament choices are a filming proposal, not actual voting or an education ranking."]}
    validation.out.mkdir(parents=True, exist_ok=True)
    (validation.out / "validation-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"package": package, **report}, ensure_ascii=False))
    return not validation.errors


def main():
    # Continue through both packages so each receives a useful diagnostic report.
    results = [validate(package) for package in ["story-v02", "worldcup-v02"]]
    if not all(results):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
