"""Check timing, cue alignment, asset references, Office render, and file integrity."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from build_package import measure

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"


def contact_sheet():
    images = sorted((OUT / "preview").glob("slide-*.png"))
    thumb_w, thumb_h, label_h, gap = 570, 321, 34, 18
    rows = (len(images) + 2) // 3
    canvas = Image.new("RGB", (3 * thumb_w + 4 * gap, rows * (thumb_h + label_h) + (rows + 1) * gap), "#E9EEF2")
    draw = ImageDraw.Draw(canvas)
    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font = ImageFont.truetype(str(font_path), 17) if font_path.exists() else ImageFont.load_default()
    for i, image_path in enumerate(images):
        x = gap + (i % 3) * (thumb_w + gap)
        y = gap + (i // 3) * (thumb_h + label_h + gap)
        with Image.open(image_path) as image:
            image = image.convert("RGB").resize((thumb_w, thumb_h), Image.Resampling.LANCZOS)
            canvas.paste(image, (x, y))
        caption = f"P{i + 1:02d} | MAIN" if i < 10 else f"A{i - 9:02d} | APPENDIX / HIDDEN"
        draw.text((x + 5, y + thumb_h + 7), caption, font=font, fill="#173449")
    canvas.save(OUT / "preview/contact-sheet.jpg", quality=92)


def main():
    film = json.loads((ROOT / "content/film.json").read_text())
    timing = json.loads((OUT / "timing-report.json").read_text())
    errors = []
    checks = []

    def check(condition, message):
        (checks if condition else errors).append(message)

    scenes = film["scenes"]
    check(len(scenes) == 10, "본편 10장면")
    check(sum(scene["duration"] for scene in scenes) == 275, "본편 275초")
    check(sum(segment["duration"] for segment in film["fixed_segments"]) == 25, "고정 영상 25초")
    check(film["total_seconds"] == 300, "전체 목표 300초")
    boundaries = sorted([(scene["start"], scene["start"] + scene["duration"]) for scene in scenes] +
                        [(segment["start"], segment["start"] + segment["duration"]) for segment in film["fixed_segments"]])
    check(boundaries[0][0] == 0 and boundaries[-1][1] == 300 and
          all(left[1] == right[0] for left, right in zip(boundaries, boundaries[1:])),
          "00:00~05:00 전 구간 연속, 공백·겹침 없음")
    check(timing["scenes"] == [measure(scene, film) for scene in scenes], "현재 대본과 발화량 보고서 일치")
    check(timing["whole_film_estimates_seconds"]["slow"] < 600, "고정 영상을 포함한 느린 발화 추정 10분 미만")
    facts = (ROOT / "docs/03-facts-and-assets.md").read_text()
    script = (ROOT / "docs/06-script.md").read_text()
    board = (ROOT / "docs/05-storyboard.md").read_text()
    prompt = (OUT / "teleprompter.txt").read_text()
    for scene in scenes:
        raw = " ".join(turn["text"] for turn in scene["turns"])
        check(raw.count(scene["cue"]["anchor"]) == 1, f"{scene['id']} 자료 등장 큐 유일")
        check(f"[{scene['slide']} 등장]" in script and f"[{scene['slide']} 등장]" in prompt,
              f"{scene['id']} 대본·프롬프터 큐 동기화")
        check(scene["slide"] in board, f"{scene['id']} 콘티 참조 존재")
        for asset in scene["assets"]:
            check(asset in facts, f"{scene['id']} {asset} 원본 수집 목록 존재")
        check(measure(scene, film)["estimates_seconds"]["slow"] <= scene["duration"],
              f"{scene['id']} 느린 속도 추정이 배정 시간 안에 들어감")

    path = OUT / "learning-fair-2026.pptx"
    with zipfile.ZipFile(path) as archive:
        check(archive.testzip() is None, "PPTX ZIP 구조 정상")
    deck = Presentation(path)
    check(len(deck.slides) == 13, "PPTX 본편 10장+부록 3장")
    check(abs(deck.slide_width / deck.slide_height - 16 / 9) < .001, "PPTX 16:9")
    for i, slide in enumerate(deck.slides):
        expected = f"P{i + 1:02d}" if i < 10 else f"A{i - 9:02d}"
        values = [shape.text for shape in slide.shapes if shape.has_text_frame]
        check(expected in values, f"PPTX {i + 1}쪽의 자료 ID {expected}")
        check(slide._element.get("show", "1") == ("1" if i < 10 else "0"), f"{expected} 본편/숨김 부록 상태")
        check(bool(slide.notes_slide.notes_text_frame.text.strip()), f"{expected} 발표자 노트 존재")
        if i < 10:
            for turn in scenes[i]["turns"]:
                # The cue annotations split a sentence, so check speech stripped of cues.
                clean_notes = re.sub(r"\[[^\]]+\] ?", "", slide.notes_slide.notes_text_frame.text)
                check(turn["text"] in clean_notes, f"{expected} 발표자 노트의 발화 동기화")
        for shape in slide.shapes:
            check(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= deck.slide_width + 20 and
                  shape.top + shape.height <= deck.slide_height + 20,
                  f"{expected} 도형 {shape.shape_id} 슬라이드 영역 안에 배치")

    pdf = fitz.open(OUT / "learning-fair-2026.pdf")
    check(len(pdf) == 13, "PowerPoint에서 실제 렌더한 PDF 13쪽")
    for i, page in enumerate(pdf):
        value = re.sub(r"\s+", "", page.get_text())
        expected = f"P{i + 1:02d}" if i < 10 else f"A{i - 9:02d}"
        check(expected in value, f"PDF {expected} 자료 번호 렌더")
        check("�" not in value, f"PDF {expected} 한글 대체문자 없음")
        if i == 5:
            check(all(word in value for word in ["로그패턴감지", "정형메시지", "메신저동보", "수작업대체", "확대방향"]), "PDF IAP 기능·현재/계획 구분 렌더")
        if i == 10:
            check("10월" in value and "예정" in value and "미개최" in value, "PDF 10월 예정 상태 렌더")
        preview = OUT / "preview" / f"slide-{i + 1:02d}.png"
        check(preview.exists(), f"{expected} 실제 PowerPoint PNG 미리보기 존재")
        if preview.exists():
            with Image.open(preview) as image:
                image.verify()
    pdf.close()
    for entry in json.loads((ROOT / "assets/references/manifest.json").read_text()):
        asset = ROOT / entry["file"]
        check(asset.exists() and hashlib.sha256(asset.read_bytes()).hexdigest() == entry["sha256"],
              f"외부 원본 {asset.name} SHA-256 일치")
    for doc in sorted((ROOT / "docs").glob("*.md")):
        check("이 문서에서 해야 할 일" in doc.read_text(), f"{doc.name} 작업 목적 명시")
    contact_sheet()
    report = {"status": "PASS" if not errors else "FAIL", "passed_checks": len(checks),
              "errors": errors, "limits": ["발화 추정이며 리허설 실측 아님", "내부 사진·회사 KV 미수급", "실제 영상 제작·편집은 별도 단계"]}
    (OUT / "validation-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
