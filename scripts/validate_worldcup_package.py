"""Validate only the additional tournament package and build its contact sheet.

Run after build_worldcup_package.py and the Office PDF/PNG export. The source
selection is a filming proposal, not a team vote or an official evaluation.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from pathlib import Path

import fitz
from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageStat
from pptx import Presentation

from build_package import FONT, cue_text, measure, name, timecode

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/worldcup"
DOCS = ROOT / "docs/worldcup"
SOURCE = ROOT / "content/worldcup-film.json"
DECK = OUT / "learning-fair-2026-worldcup.pptx"
PDF = OUT / "learning-fair-2026-worldcup.pdf"
PREVIEW = OUT / "preview"
CHOICE_DECK = OUT / "choice-states.pptx"
CHOICE_PDF = OUT / "choice-states.pdf"
CHOICE_PREVIEW = OUT / "choice-preview"
PAGE_IDS = [f"P{i:02d}" for i in range(1, 11)] + [f"A{i:02d}" for i in range(1, 4)]
ALLOWED_FONTS = {FONT, "Malgun Gothic", "맑은 고딕"}
EXPECTED_CANDIDATES = {
    "seminar": "AI 4 AI 세미나", "job-fair": "Job Fair",
    "master-tour": "명장 투어", "tech-x": "Tech X",
}
EXPECTED_MATCHES = [
    ("M01", "준결승 1", ["seminar", "job-fair"], "job-fair"),
    ("M02", "준결승 2", ["master-tour", "tech-x"], "tech-x"),
    ("M03", "결승", ["job-fair", "tech-x"], "tech-x"),
]


def compact(value: str) -> str:
    return re.sub(r"\s+", "", value)


class Validation:
    def __init__(self):
        self.passed = []
        self.errors = []
        self.unavailable = []
        self.metrics = {}

    def check(self, condition, message):
        (self.passed if condition else self.errors).append(message)

    def require(self, path: Path) -> bool:
        if path.is_file():
            self.check(path.stat().st_size > 0, f"{path.relative_to(ROOT)} 비어 있지 않음")
            return True
        self.unavailable.append(f"미생성 또는 미수급: {path.relative_to(ROOT)}")
        return False

    def text(self, path: Path):
        if not self.require(path):
            return None
        try:
            return path.read_text(encoding="utf-8-sig")
        except (OSError, UnicodeError) as error:
            self.errors.append(f"{path.name} 텍스트 열람 실패: {error}")
            return None

    def json(self, path: Path):
        value = self.text(path)
        if value is None:
            return None
        try:
            return json.loads(value)
        except json.JSONDecodeError as error:
            self.errors.append(f"{path.name} JSON 형식 오류: {error}")
            return None

    def section(self, label, action):
        try:
            return action()
        except Exception as error:
            self.errors.append(f"{label} 검증 중 오류: {type(error).__name__}: {error}")
            return None


def check_source(film, validation):
    check = validation.check
    scenes = film["scenes"]
    fixed = film["fixed_segments"]
    check([s["id"] for s in scenes] == [f"S{i:02d}" for i in range(1, 11)], "본편 S01~S10 순서·중복 없음")
    check([s["slide"] for s in scenes] == PAGE_IDS[:10], "본편 P01~P10 자료 대응")
    check([s["id"] for s in film["appendix"]] == PAGE_IDS[10:], "부록 A01~A03 자료 대응")
    check(sum(s["duration"] for s in scenes) == 275, "본편 배정 275초")
    check(film["total_seconds"] == 300, "전체 목표 300초")
    actual_fixed = sorted((s["start"], s["duration"]) for s in fixed)
    check(actual_fixed == [(0, 5), (5, 15), (295, 5)], "회사 인트로 5초·AI 소개 15초·아웃트로 5초")
    boundaries = sorted((s["start"], s["start"] + s["duration"]) for s in scenes + fixed)
    check(all(start < end for start, end in boundaries), "모든 구간의 길이가 양수")
    check(boundaries[0][0] == 0 and boundaries[-1][1] == 300 and
          all(left[1] == right[0] for left, right in zip(boundaries, boundaries[1:])),
          "00:00~05:00 전 구간 연속·공백·겹침 없음")
    check(scenes[0]["start"] == 20 and scenes[-1]["start"] + scenes[-1]["duration"] == 295,
          "스튜디오 본편 00:20~04:55")
    check(set(film["reading_rates"]) == {"slow", "baseline", "fast"} and
          all(isinstance(rate, (int, float)) and rate > 0 for rate in film["reading_rates"].values()),
          "느림·기준·빠름의 발화 속도 유효")
    check(film["reading_rates"]["slow"] <= film["reading_rates"]["baseline"] <= film["reading_rates"]["fast"],
          "발화 속도 가정 순서")
    iap = [s for s in scenes if s["id"] == "S06"]
    check(len(iap) == 1 and iap[0]["slide"] == "P06" and iap[0]["duration"] == 40,
          "IAP 대표 사례는 S06/P06 한 장·40초")
    if iap:
        check({"IMG08", "IMG09", "IMG10"}.issubset(iap[0]["assets"]),
              "S06에 IAP 기능·6월 소개·7월 실습 원본 모두 대응")
    iap_scenes = [s["id"] for s in scenes
                  if re.search(r"(?<![A-Za-z])IAP(?![A-Za-z])|Infra Agent Platform",
                               " ".join(turn["text"] for turn in s["turns"]))]
    check(iap_scenes == ["S06"], "IAP 발화가 다른 본편 장면으로 확장되지 않음")
    validation.metrics["main_allocated_seconds"] = sum(s["duration"] for s in scenes)
    validation.metrics["total_seconds"] = film["total_seconds"]


def check_tournament(film, validation):
    check = validation.check
    tournament = film["tournament"]
    check(tournament["status"] == "촬영용 선택안", "토너먼트는 촬영용 선택안")
    check(tournament.get("actual_vote") is False, "실제 팀원 투표 결과로 표시하지 않음")
    check(bool(tournament["criterion"].strip()), "공통 선택 기준 존재")
    roster = tournament["candidates"]
    check(len(roster) == 4 and {c["id"]: c["name"] for c in roster} == EXPECTED_CANDIDATES,
          "네 후보의 원본 ID·명칭 일치")
    actual = [(m["id"], m["round"], m["candidates"], m["winner"]) for m in tournament["matches"]]
    check(actual == EXPECTED_MATCHES, "4강 Job Fair/Tech X 진출·결승 Tech X 선택안")
    for match in tournament["matches"]:
        check(len(match["candidates"]) == 2 and len(set(match["candidates"])) == 2 and
              all(candidate in EXPECTED_CANDIDATES for candidate in match["candidates"]) and
              match["winner"] in match["candidates"], f"{match['id']} 후보·선택 결과 유효")
    matches = tournament["matches"]
    if len(matches) == 3:
        check(matches[2]["candidates"] == [matches[0]["winner"], matches[1]["winner"]],
              "결승 후보가 준결승 선택과 연결")
    disclaimer = tournament["disclaimer"]
    excludes_ranking = ("공식" in disclaimer and "평가" in disclaimer) or ("교육 효과" in disclaimer and "순위" in disclaimer)
    check(all(word in disclaimer for word in ["촬영", "선택", "실제", "투표"]) and excludes_ranking and
          bool(re.search(r"아니|않|무관|미실시", disclaimer)), "선택안·실제 투표 및 교육 평가 순위 아님을 명시")
    validation.metrics["selection_status"] = tournament["status"]
    validation.metrics["actual_vote"] = tournament.get("actual_vote")


def markdown_scene_blocks(value):
    headings = list(re.finditer(r"(?m)^#{2,3}[^\n]*\b(S\d{2})\b[^\n]*$", value))
    return {match.group(1): value[match.end():headings[i + 1].start() if i + 1 < len(headings) else len(value)]
            for i, match in enumerate(headings)}


def prompter_scene_blocks(value):
    headings = list(re.finditer(r"(?m)^\[(S\d{2})\s*/\s*(P\d{2})\s*/[^\n]*\]\s*$", value))
    return {match.group(1): (match.group(2), value[match.end():headings[i + 1].start() if i + 1 < len(headings) else len(value)])
            for i, match in enumerate(headings)}


def prompter_turns(value):
    headings = list(re.finditer(r"(?m)^<([^>\n]+)>\s*$", value))
    return [(match.group(1), value[match.end():headings[i + 1].start() if i + 1 < len(headings) else len(value)].strip())
            for i, match in enumerate(headings)]


def check_documents(film, timing, validation):
    check = validation.check
    scenes = film["scenes"]
    measures = [measure(scene, film) for scene in scenes]
    validation.metrics["total_spoken_syllables"] = sum(row["spoken_syllables"] for row in measures)
    for scene, row in zip(scenes, measures, strict=True):
        raw = " ".join(turn["text"] for turn in scene["turns"])
        check(isinstance(scene["pause_seconds"], (int, float)) and scene["pause_seconds"] >= 0,
              f"{scene['id']} 호흡 예산 유효")
        check(row["estimates_seconds"]["slow"] <= scene["duration"], f"{scene['id']} 느린 발화 추정이 배정 안에 들어감")
        for cue in [scene["cue"], *scene.get("secondary_cues", [])]:
            check(bool(cue["anchor"]) and raw.count(cue["anchor"]) == 1,
                  f"{scene['id']} 자료 큐 ‘{cue['anchor']}’ 발화 안에서 유일")
        for turn in scene["turns"]:
            check(turn["speaker"] in film["presenters"], f"{scene['id']} 발화자 이름 대응")
    totals = {key: round(sum(row["estimates_seconds"][key] for row in measures), 2)
              for key in film["reading_rates"]}
    whole = {key: round(value + 25, 2) for key, value in totals.items()}
    validation.metrics["main_estimates_seconds"] = totals
    check(whole["slow"] < 600, "고정 영상 포함 느린 발화 추정 10분 미만")
    if timing is not None:
        check(timing["scenes"] == measures, "현재 대본과 장면별 발화량 보고서 일치")
        check(timing["main_estimates_seconds"] == totals and timing["whole_film_estimates_seconds"] == whole,
              "본편·전체 발화 추정 합계 일치")
        check(timing["total_spoken_syllables"] == sum(row["spoken_syllables"] for row in measures), "총 발음 환산 음절 수 일치")
        check(timing["reading_rates_syllables_per_second"] == film["reading_rates"], "속도 가정 보고서와 원본 일치")
        check(timing["main_allocated_seconds"] == 275 and timing["fixed_seconds"] == 25 and
              timing["iap_allocated_seconds"] == 40, "시간 보고서의 275+25초·IAP40초 일치")
        check(timing["version"] == film["version"] and timing["date"] == film["date"], "시간 보고서 원본 버전·날짜 일치")
    script = validation.text(DOCS / "02-script.md")
    board = validation.text(DOCS / "01-storyboard.md")
    prompt = validation.text(OUT / "teleprompter.txt")
    facts = validation.text(ROOT / "docs/03-facts-and-assets.md")
    registered = set(re.findall(r"(?m)^\|\s*(IMG\d{2})\s*\|", facts or ""))
    for label, value in [("대본", script), ("콘티", board)]:
        if value is not None:
            check("이 문서에서 해야 할 일" in value, f"월드컵 {label}의 작업 목적 명시")
            check(film["tournament"]["disclaimer"] in value, f"월드컵 {label}의 선택안 성격 원본과 일치")
            check(film["tournament"]["criterion"] in value, f"월드컵 {label}의 공통 선택 기준 일치")
            check(set(re.findall(r"\bIMG\d{2}\b", value)).issubset(registered), f"월드컵 {label} 내부 원본 ID 등록")
    script_blocks = markdown_scene_blocks(script) if script is not None else {}
    board_blocks = markdown_scene_blocks(board) if board is not None else {}
    prompt_blocks = prompter_scene_blocks(prompt) if prompt is not None else {}
    for scene in scenes:
        sid, slide = scene["id"], scene["slide"]
        for asset in scene["assets"]:
            check(asset in registered, f"{sid} {asset} 내부 원본 수집 목록 대응")
        if script is not None:
            block = script_blocks.get(sid, "")
            check(bool(block), f"{sid} 대본 장면 구분 존재")
            for turn in scene["turns"]:
                expected = f"{name(turn['speaker'], film)}: {cue_text(turn, scene, film)}"
                check(compact(expected) in compact(block.replace("**", "")), f"{sid} 대본의 화자·발화·자료 큐 동기화")
            check(f"[{slide} 등장]" in block, f"{sid} 대본의 주 자료 큐")
        if board is not None:
            block = board_blocks.get(sid, "")
            check(bool(block) and slide in block and scene["cue"]["anchor"] in block, f"{sid} 콘티의 자료·발화 큐 대응")
            check(scene["camera"] in block and scene["hold"] in block, f"{sid} 콘티의 촬영·편집 지시 일치")
        if prompt is not None:
            actual_slide, block = prompt_blocks.get(sid, ("", ""))
            check(actual_slide == slide, f"{sid} 프롬프터 장면·자료 ID 대응")
            actual = [(speaker, compact(speech)) for speaker, speech in prompter_turns(block)]
            expected = [(name(turn["speaker"], film), compact(cue_text(turn, scene, film, phonetic=True)))
                        for turn in scene["turns"]]
            check(actual == expected, f"{sid} 프롬프터 화자·발음·자료 큐 정확히 일치")
            spoken_only = " ".join(re.sub(r"\[[^\]]*\]", "", speech) for _, speech in actual)
            check(not re.search(r"[A-Za-z0-9]", spoken_only), f"{sid} 읽는 문장에 미치환 영문·숫자 없음")
    rehearsal = OUT / "rehearsal.csv"
    if validation.require(rehearsal):
        with rehearsal.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        check([row["scene"] for row in rows] == [scene["id"] for scene in scenes], "리허설 CSV의 본편 10장면 순서")
        for scene, measured, estimate in zip(scenes, rows, measures):
            check(measured["PPT"] == scene["slide"] and measured["start"] == timecode(scene["start"]) and
                  measured["end"] == timecode(scene["start"] + scene["duration"]), f"{scene['id']} 리허설 자료·타임코드 대응")
            check(float(measured["allocated_seconds"]) == scene["duration"] and
                  float(measured["syllables"]) == estimate["spoken_syllables"] and
                  math.isclose(float(measured["estimated_baseline_seconds"]), estimate["estimates_seconds"]["baseline"]),
                  f"{scene['id']} 리허설 배정·발화량·추정 대응")
            value = measured["measured_seconds"].strip()
            check(not value or (math.isfinite(float(value)) and float(value) >= 0), f"{scene['id']} 리허설 실측 칸 유효")
        validation.metrics["measured_rehearsal_scenes"] = sum(bool(row["measured_seconds"].strip()) for row in rows)


def slide_strings(slide):
    return [shape.text for shape in slide.shapes if shape.has_text_frame]


def named_text(slide, shape_name):
    return [shape.text for shape in slide.shapes if shape.name == shape_name and shape.has_text_frame]


def check_deck(film, validation):
    check = validation.check
    if not validation.require(DECK):
        return None
    with zipfile.ZipFile(DECK) as archive:
        check(archive.testzip() is None, "월드컵 PPTX ZIP 구조 정상")
        animations = any(b"<p:timing" in archive.read(item) for item in archive.namelist()
                         if re.fullmatch(r"ppt/slides/slide\d+\.xml", item))
        check(not animations, "월드컵 PPTX는 정적 자료·모션은 편집 지시")
    deck = Presentation(DECK)
    check(len(deck.slides) == 13, "월드컵 PPTX 본편10+부록3장")
    check(abs(deck.slide_width / deck.slide_height - 16 / 9) < .001, "월드컵 PPTX 16:9")
    check(sum(s._element.get("show", "1") != "0" for s in deck.slides) == 10, "PPTX 표시 본편 10장")
    check(sum(s._element.get("show", "1") == "0" for s in deck.slides) == 3, "PPTX 숨김 부록 3장")
    for i, slide in enumerate(deck.slides):
        expected = PAGE_IDS[i] if i < len(PAGE_IDS) else f"EXTRA-{i + 1}"
        values = slide_strings(slide)
        check(expected in values, f"{expected} PPTX 자료 ID 존재")
        check(slide._element.get("show", "1") == ("1" if i < 10 else "0"), f"{expected} 표시·숨김 상태")
        notes = slide.notes_slide.notes_text_frame.text
        check(bool(notes.strip()), f"{expected} 발표자 노트 존재")
        if i < min(10, len(film["scenes"])):
            scene = film["scenes"][i]
            check(film["tournament"]["disclaimer"] in notes, f"{expected} 노트에 촬영용 선택안 명시")
            check(scene["camera"] in notes and scene["hold"] in notes, f"{expected} 노트 촬영·편집 지시 동기화")
            for turn in scene["turns"]:
                expected_turn = f"{name(turn['speaker'], film)}: {cue_text(turn, scene, film)}"
                check(compact(expected_turn) in compact(notes), f"{expected} 노트 화자·발화·자료 큐 동기화")
            for cue in scene.get("secondary_cues", []):
                check(cue["anchor"] in notes and cue["label"] in notes and cue["effect"] in notes,
                      f"{expected} 노트의 추가 자료 큐 동기화")
        for shape in slide.shapes:
            check(shape.left >= 0 and shape.top >= 0 and shape.width >= 0 and shape.height >= 0 and
                  shape.left + shape.width <= deck.slide_width + 20 and
                  shape.top + shape.height <= deck.slide_height + 20,
                  f"{expected} 도형 {shape.shape_id} 슬라이드 영역 안에 배치")
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue
                        check(run.font.name in ALLOWED_FONTS and run.font.size is not None and run.font.size > 0,
                              f"{expected} 도형 {shape.shape_id} 명시적 폰트·크기 유효")
                        east_asia = run._r.xpath("./a:rPr/a:ea")
                        check(bool(east_asia) and east_asia[0].get("typeface") in ALLOWED_FONTS,
                              f"{expected} 도형 {shape.shape_id} 한글 폰트 지정")
        check(set(re.findall(r"\bIMG\d{2}\b", " ".join(values) + notes)).issubset({f"IMG{x:02d}" for x in range(1, 15)}),
              f"{expected} PPTX 내부 원본 ID 유효")
    if len(deck.slides) >= 10:
        check_selection_render(deck, film, validation)
    return deck


def check_selection_render(deck, film, validation):
    check = validation.check
    tournament = film["tournament"]
    candidate_names = {c["id"]: c["name"] for c in tournament["candidates"]}
    p02 = deck.slides[1]
    text_p02 = compact(" ".join(slide_strings(p02)))
    check(all(compact(value) in text_p02 for value in candidate_names.values()), "P02에 네 후보 렌더")
    for shape_name, value in [("bracket-final-left", "준결승 1의 선택"), ("bracket-final-right", "준결승 2의 선택")]:
        actual = [shape.text for shape in p02.shapes if shape.name == shape_name and shape.has_text_frame]
        check(actual == [value], f"P02 {shape_name} 결승 자리 미공개")
    for index, match_index in [(3, 0), (6, 1)]:
        winner = candidate_names[tournament["matches"][match_index]["winner"]]
        check(compact(f"결승 진출 · {winner}") in compact(" ".join(slide_strings(deck.slides[index]))),
              f"P{index + 1:02d} 준결승 진출자 원본 메타와 일치")
    p08 = deck.slides[7]
    values = compact(" ".join(slide_strings(p08)))
    final = tournament["matches"][2]
    winner = candidate_names[final["winner"]]
    check(named_text(p08, "bracket-winner") == [winner], "P08 최종 Tech X 추천 렌더")
    for shape_name, candidate_id in zip(["bracket-final-left", "bracket-final-right"], final["candidates"], strict=True):
        actual = [shape.text for shape in p08.shapes if shape.name == shape_name and shape.has_text_frame]
        check(actual == [candidate_names[candidate_id]], f"P08 {shape_name} 결승 대진 원본과 일치")
    check("촬영용" in values and ("제안" in values or "선택안" in values), "P08 화면의 추천은 촬영용 선택안으로 표시")
    iap_text = compact(" ".join(slide_strings(deck.slides[5])))
    check(all(word in iap_text for word in ["IAP", "로그패턴감지", "정형메시지", "메신저동보", "수작업대체", "확대목표", "6월", "7월", "25명"]),
          "P06 IAP 기능·현재/목표·소개/실습 한 장에 렌더")


def check_choice_states(film, main_deck, validation):
    check = validation.check
    available_deck = validation.require(CHOICE_DECK)
    available_pdf = validation.require(CHOICE_PDF)
    png_paths = [CHOICE_PREVIEW / f"slide-{i:02d}.png" for i in [1, 2]]
    available_png = [validation.require(path) for path in png_paths]
    actual_names = {path.name for path in CHOICE_PREVIEW.glob("slide-*.png")}
    expected_names = {path.name for path in png_paths}
    check(not actual_names - expected_names, "결승 PNG에 불필요한 장면 없음")
    if all(available_png):
        check(actual_names == expected_names, "결승 선택 전·후 PNG 두 장 존재")
    scene = next(s for s in film["scenes"] if s["id"] == "S08")
    check("선택전" in compact(scene["cue"]["effect"]) and "choice-preview/slide-01.png" in scene["cue"]["effect"],
          "S08 최초 자료 큐는 선택 전 PNG")
    reveal_cues = [cue for cue in scene.get("secondary_cues", [])
                   if compact(cue["anchor"]) == compact("저는 Tech X입니다")]
    check(len(reveal_cues) == 1 and "선택후" in compact(reveal_cues[0]["label"] + " " + reveal_cues[0]["effect"]) and
          "choice-preview/slide-02.png" in reveal_cues[0]["effect"], "S08 선택 발화에서 선택 후 PNG로 전환")
    states = None
    if available_deck:
        with zipfile.ZipFile(CHOICE_DECK) as archive:
            check(archive.testzip() is None, "선택 상태 PPTX ZIP 구조 정상")
        states = Presentation(CHOICE_DECK)
        check(len(states.slides) == 2, "선택 상태 PPTX 선택 전·후 2장")
        check(abs(states.slide_width / states.slide_height - 16 / 9) < .001, "선택 상태 PPTX 16:9")
        for i, slide in enumerate(states.slides):
            check("P08" in slide_strings(slide), f"선택 상태 {i + 1} 자료 ID P08")
            check(slide._element.get("show", "1") == "1", f"선택 상태 {i + 1} 표시 슬라이드")
            check(bool(slide.notes_slide.notes_text_frame.text.strip()), f"선택 상태 {i + 1} 편집 큐 노트 존재")
        if len(states.slides) == 2:
            before, after = states.slides
            check(named_text(before, "bracket-winner") == ["오늘의 추천"], "선택 전 우승자 미공개")
            check(named_text(after, "bracket-winner") == ["Tech X"], "선택 후 우승자 Tech X 공개")
            check(not named_text(before, "bracket-winner-reason") and bool(named_text(after, "bracket-winner-reason")),
                  "선택 이유는 선택 후에만 표시")
            for state in [before, after]:
                check(named_text(state, "bracket-final-left") == ["Job Fair"] and
                      named_text(state, "bracket-final-right") == ["Tech X"], "전·후 화면의 결승 대진 동일")
            if main_deck is not None and len(main_deck.slides) >= 8:
                check(slide_strings(main_deck.slides[7]) == slide_strings(after), "주 PPTX P08은 선택 후 완성 상태와 동일")
    if available_pdf:
        with fitz.open(CHOICE_PDF) as pdf:
            check(len(pdf) == 2, "선택 상태 PDF 실제 열람 2쪽")
            for i, page in enumerate(pdf):
                value = compact(page.get_text())
                check("P08" in value and "�" not in value, f"선택 상태 PDF {i + 1} 자료 ID·문자 정상")
                check(abs(page.rect.width / page.rect.height - 16 / 9) < .001, f"선택 상태 PDF {i + 1} 16:9")
                if states is not None and i < len(states.slides):
                    check(all(compact(text) in value for text in slide_strings(states.slides[i])),
                          f"선택 상태 PDF {i + 1} PPTX 텍스트 일치")
                if i == 0:
                    check("오늘의추천" in value and "새로운일에함께참여" not in value, "선택 전 PDF에 결과·선택 이유 조기 노출 없음")
                elif i == 1:
                    check("새로운일에함께참여" in value, "선택 후 PDF에 선택 이유 공개")
                if i < 2 and available_png[i]:
                    with Image.open(png_paths[i]) as png:
                        png.load()
                        check(png.format == "PNG" and abs(png.width / png.height - 16 / 9) < .01,
                              f"선택 상태 PNG {i + 1} 정상·16:9")
                        pix = page.get_pixmap(matrix=fitz.Matrix(.2, .2), alpha=False)
                        rendered = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                        a = png.convert("L").resize((96, 54), Image.Resampling.LANCZOS)
                        b = rendered.convert("L").resize((96, 54), Image.Resampling.LANCZOS)
                        check(ImageStat.Stat(ImageChops.difference(a, b)).mean[0] / 255 <= .12,
                              f"선택 상태 PNG {i + 1} PDF 화면과 일치")
    if all(available_png):
        with Image.open(png_paths[0]) as before, Image.open(png_paths[1]) as after:
            check(before.size == after.size and
                  ImageChops.difference(before.convert("RGB"), after.convert("RGB")).getbbox() is not None,
                  "선택 전·후 실제 PNG의 화면 픽셀이 서로 다름")


def check_render(deck, validation):
    check = validation.check
    available = [validation.require(PREVIEW / f"slide-{i:02d}.png") for i in range(1, 14)]
    actual_names = {path.name for path in PREVIEW.glob("slide-*.png")}
    expected_names = {f"slide-{i:02d}.png" for i in range(1, 14)}
    check(not actual_names - expected_names, "PNG 미리보기에 불필요한 장면 없음")
    if all(available):
        check(actual_names == expected_names, "PNG 미리보기 01~13 존재")
    if not validation.require(PDF):
        return
    with fitz.open(PDF) as pdf:
        check(len(pdf) == 13, "월드컵 PDF 실제 열람 13쪽")
        fonts = set()
        differences = {}
        for i, page in enumerate(pdf):
            expected = PAGE_IDS[i] if i < len(PAGE_IDS) else f"EXTRA-{i + 1}"
            extracted = page.get_text()
            page_text = compact(extracted)
            check(expected in page_text, f"PDF {expected} 자료 ID 렌더")
            check("�" not in extracted, f"PDF {expected} 대체문자 없음")
            check(abs(page.rect.width / page.rect.height - 16 / 9) < .001, f"PDF {expected} 16:9")
            fonts.update(font[3] for font in page.get_fonts())
            if deck is not None and i < len(deck.slides):
                for value in slide_strings(deck.slides[i]):
                    check(compact(value) in page_text, f"PDF {expected} PPTX 텍스트 렌더 일치: {value[:30]}")
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        if not span["text"].strip():
                            continue
                        x0, y0, x1, y1 = span["bbox"]
                        check(x0 >= -1 and y0 >= -1 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1,
                              f"PDF {expected} 텍스트 슬라이드 밖 넘침 없음")
            if i < 13 and available[i]:
                with Image.open(PREVIEW / f"slide-{i + 1:02d}.png") as png:
                    png.load()
                    check(png.format == "PNG" and png.width > 0 and png.height > 0, f"{expected} PNG 이미지 정상")
                    check(abs(png.width / png.height - 16 / 9) < .01, f"{expected} PNG 16:9")
                    check(any(low != high for low, high in png.convert("RGB").getextrema()), f"{expected} PNG 빈 단색 이미지 아님")
                    # Low-resolution luminance comparison tolerates antialiasing
                    # differences between the Office PNG and PDF renderers.
                    pix = page.get_pixmap(matrix=fitz.Matrix(.2, .2), alpha=False)
                    rendered = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    size = (96, 54)
                    a = png.convert("L").resize(size, Image.Resampling.LANCZOS)
                    b = rendered.convert("L").resize(size, Image.Resampling.LANCZOS)
                    difference = ImageStat.Stat(ImageChops.difference(a, b)).mean[0] / 255
                    differences[expected] = round(difference, 4)
                    check(difference <= .12, f"{expected} PNG·PDF 화면 내용 일치(축소 명도 차이)")
            if i == 10:
                check(all(word in page_text for word in ["10월", "예정", "미개최", "OutSystems"]), "A01 10월 세미나 미개최·예정 렌더")
            if i == 11:
                check("IMG08" in page_text and "IMG01" in page_text and "원본" in page_text, "A02 촬영 원본 수집 자료 렌더")
            if i == 12:
                check(all(word in page_text for word in ["실제투표결과아님", "모션", "40초"]), "A03 선택변경·모션·IAP40초 확인 부록 렌더")
        check(any("Malgun" in font or "맑은" in font for font in fonts), "실제 PDF에 맑은 고딕 계열 렌더")
        validation.metrics["pdf_font_names"] = sorted(fonts)
        validation.metrics["preview_pdf_luminance_difference"] = differences


def contact_sheet(validation):
    paths = [PREVIEW / f"slide-{i:02d}.png" for i in range(1, 14)]
    if not all(path.is_file() for path in paths):
        return
    thumb_w, thumb_h, label_h, gap = 570, 321, 34, 18
    rows = (len(paths) + 2) // 3
    canvas = Image.new("RGB", (3 * thumb_w + 4 * gap, rows * (thumb_h + label_h) + (rows + 1) * gap), "#E9EEF2")
    draw = ImageDraw.Draw(canvas)
    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font = ImageFont.truetype(str(font_path), 17) if font_path.exists() else ImageFont.load_default()
    for i, path in enumerate(paths):
        x, y = gap + (i % 3) * (thumb_w + gap), gap + (i // 3) * (thumb_h + label_h + gap)
        with Image.open(path) as image:
            canvas.paste(image.convert("RGB").resize((thumb_w, thumb_h), Image.Resampling.LANCZOS), (x, y))
        caption = f"WORLD CUP / {PAGE_IDS[i]} | " + ("MAIN" if i < 10 else "APPENDIX / HIDDEN")
        draw.text((x + 5, y + thumb_h + 7), caption, font=font, fill="#173449")
    target = PREVIEW / "contact-sheet.jpg"
    canvas.save(target, quality=92)
    validation.check(target.is_file(), "월드컵 13장 컨택트시트 생성")


def main():
    validation = Validation()
    film = validation.json(SOURCE)
    timing = validation.json(OUT / "timing-report.json")
    deck = None
    if film is not None:
        validation.section("원본 시간·IAP 범위", lambda: check_source(film, validation))
        validation.section("토너먼트 메타데이터", lambda: check_tournament(film, validation))
        validation.section("발화량·대본·콘티·프롬프터·리허설", lambda: check_documents(film, timing, validation))
        deck = validation.section("PPTX·노트·도형·폰트", lambda: check_deck(film, validation))
        validation.section("결승 선택 전·후 자료", lambda: check_choice_states(film, deck, validation))
    validation.section("PDF·PNG 렌더", lambda: check_render(deck, validation))
    validation.section("컨택트시트", lambda: contact_sheet(validation))
    artifacts = [SOURCE, DECK, PDF, CHOICE_DECK, CHOICE_PDF, DOCS / "01-storyboard.md", DOCS / "02-script.md",
                 OUT / "teleprompter.txt", OUT / "rehearsal.csv", OUT / "timing-report.json"]
    status = "FAIL" if validation.errors else "UNAVAILABLE" if validation.unavailable else "PASS"
    report = {
        "status": status, "package": "worldcup", "passed_checks": len(validation.passed),
        "errors": validation.errors, "unavailable": validation.unavailable,
        "metrics": validation.metrics,
        "artifact_sha256": {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest()
                            for path in artifacts if path.is_file()},
        "limits": [
            "발화량은 계획 추정이며 실제 리허설 기록을 대체하지 않음",
            "선택 결과는 촬영용 선택안이며 실제 투표·공식 팀원 평가 아님",
            "내부 사진·회사 KV는 미수급, 실제 촬영·편집은 별도",
            "도형·PDF 텍스트 영역과 파일 일치를 자동 검증하되 최종 가독성·모션은 수동 검토 필요",
        ],
    }
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "validation-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    if status != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
