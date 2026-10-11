"""Build the standalone balance-game package from its JSON source."""

from __future__ import annotations

import copy
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import presentation_common as drawing
import production_common as production
from build_story_pair import draw_slot
from presentation_common import box, connector, text

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = "balance-v03"
SOURCE = ROOT / "content" / f"{PACKAGE}-film.json"
OUT = ROOT / "deliverables" / PACKAGE
DOCS = ROOT / "docs" / PACKAGE
CHOICE_STATES = [("PRE", None), ("POST-A", "A"), ("POST-B", "B")]
LAYOUTS = {"opening", "balance", "paired_evidence", "experience", "iap", "invitation", "closing"}


def prepare_film(source):
    film = copy.deepcopy(source)
    film.setdefault("display_name", "밸런스 게임형")
    film.setdefault("timeline_summary", "전체 300초 = 회사 인트로 5초 + 스튜디오 오프닝 20초 + AI 팀 소개 15초 + 나머지 스튜디오 255초 + 회사 아웃트로 5초.")
    film.setdefault("timeline_note", "스튜디오는 00:05에 시작하고, 오프닝 뒤 00:25–00:40에 별도 AI 팀 소개 영상을 삽입한다. 회사 인트로·AI 소개·아웃트로에는 스튜디오 발화가 없다. 선택·웃음·자료 전환·호흡은 각 장면 시간에 포함한다.")
    film.setdefault("build_command", ".venv/bin/python scripts/build_balance.py")
    film.setdefault("rehearsal_note", "리허설 CSV의 measured_seconds에 실측을 기록한다. 대사·화자·큐·발음·배정 시간이 바뀌면 이전 실측을 비우고 다시 측정한다. 밸런스 게임의 선택 대사는 촬영 제안이며, 실제 선택에 맞는 응답과 화면을 같은 장면 시간 안에서 사용한다.")
    film.setdefault("production_notes", [
        "PPTX는 정적 자료다. 등장·선택 강조와 자료 전환은 영상팀 편집 지시이며 장면 시간 안에서 구현한다. 숨김 부록 A01–A03은 본편에서 제외한다.",
        "선택은 프로그램의 순위나 승패를 정하지 않는다. PRE / POST-A / POST-B는 같은 질문의 공개 상태다. 실제 선택에 맞는 대체 응답을 사용하고, 촬영 후 본문 대사·공개 큐·선택 화면을 함께 확정한다.",
        "실제 사내 사진과 화면은 사용자가 넣는다. 기본 PPTX를 복사한 별도 제출 사본에 붙이고, 범용 외부 사진을 실제 팀 행사로 사용하지 않는다. 원본 번호는 docs/03-facts-and-assets.md를 따른다.",
    ])
    return film


def check_source(film):
    production.check_source(film)
    assert film["variant"] == "balance"
    assert film["official_expansion"] == "Advanced Infrastructure for Artificial Intelligence"
    assert sorted((s["start"], s["duration"]) for s in film["fixed_segments"]) == [(0, 5), (25, 15), (295, 5)]
    rounds = film["balance"]["rounds"]
    assert [item["id"] for item in rounds] == ["R01", "R02", "R03"]
    scenes = {s["id"]: s for s in film["scenes"]}
    for scene in scenes.values():
        visual = scene["visual"]
        assert visual["layout"] in LAYOUTS, scene["id"]
        if visual["layout"] in {"paired_evidence", "experience"}:
            assert len(scene.get("asset_slots", [])) == 2, scene["id"]
        if visual["layout"] in {"balance", "paired_evidence"}:
            assert len(visual["items"]) == 2, scene["id"]
        if visual["layout"] == "iap":
            assert len(visual["items"]) == 3, scene["id"]
    for item in rounds:
        assert [choice["id"] for choice in item["choices"]] == ["A", "B"], item["id"]
        assert item["selected"] in {"A", "B"}, item["id"]
        assert scenes[item["scene"]]["visual"]["layout"] == "balance", item["id"]
        assert scenes[item["scene"]]["visual"]["round"] == item["id"], item["id"]
        raw = " ".join(turn["text"] for turn in scenes[item["reveal_scene"]]["turns"])
        assert raw.count(item["reveal_anchor"]) == 1, item["id"]


def make_base(prs, scene, film, pid=None):
    visual = scene["visual"]
    stages = drawing.STAGES
    try:
        drawing.STAGES = film["stages"]
        slide = drawing.base(prs, pid or scene["slide"], visual["eyebrow"], visual["headline"], visual.get("stage"))
    finally:
        drawing.STAGES = stages
    if pid:
        marker = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == pid)
        marker.left, marker.width = Inches(10.55), Inches(2.03)
    return slide


def draw_visual(prs, scene, film, pid=None, selected=None):
    visual = scene["visual"]
    slide = make_base(prs, scene, film, pid)
    layout = visual["layout"]
    page = pid or scene["slide"]
    slots = scene.get("asset_slots", [])
    if layout == "opening":
        title = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == visual["headline"])
        for paragraph in title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(56)
        text(slide, visual["bottom"], .75, 3.14, 11.83, .8, 34, "mint", True)
        text(slide, visual["body"], .75, 4.51, 11.83, .75, 25)
        text(slide, visual["detail"], .75, 5.76, 11.83, .45, 18, "muted")
    elif layout == "balance":
        question = next(item for item in film["balance"]["rounds"] if item["id"] == visual["round"])
        for index, item in enumerate(visual["items"]):
            choice = question["choices"][index]
            x = .75 + index * 6.06
            color = "mint" if index == 0 else "lilac"
            box(slide, x, 3.2, 5.77, 2.56, line=color if selected == choice["id"] else "line", radius=True)
            text(slide, item["label"], x + .25, 3.42, 5.27, .4, 20, color, True)
            text(slide, item["value"], x + .25, 4.08, 5.27, .79, 27, bold=True)
            text(slide, item.get("detail", ""), x + .25, 5.11, 5.27, .43, 16, "muted")
        if selected:
            choice = next(item for item in question["choices"] if item["id"] == selected)
            bottom = f"지금의 선택  {selected} · {choice['name']}"
        else:
            bottom = visual.get("bottom", "지금 더 끌리는 쪽을 골라주세요")
        text(slide, bottom, .75, 6.03, 11.83, .4, 20, "mint" if selected else "muted", bool(selected))
    elif layout == "paired_evidence":
        for index, (item, slot) in enumerate(zip(visual["items"], slots, strict=True)):
            x = .75 + index * 6.06
            color = "mint" if index == 0 else "lilac"
            text(slide, item["label"], x, 3.03, 5.77, .4, 23, color, True)
            text(slide, item["value"], x, 3.52, 5.77, .39, 19, bold=True)
            text(slide, item.get("detail", ""), x, 3.98, 5.77, .35, 15, "muted")
            draw_slot(slide, page, slot, x, 4.42, 5.77, 1.5)
        text(slide, visual["bottom"], .75, 6.08, 11.83, .41, 21, "mint", True)
    elif layout == "experience":
        text(slide, visual["body"], .75, 3.02, 11.83, .6, 22, "mint", True)
        for index, slot in enumerate(slots):
            draw_slot(slide, page, slot, .75 + index * 6.06, 3.84, 5.77, 1.86)
        text(slide, visual["bottom"], .75, 5.99, 11.83, .44, 21, bold=True)
    elif layout == "iap":
        width, gap = 3.67, .41
        for index, item in enumerate(visual["items"]):
            x = .75 + index * (width + gap)
            box(slide, x, 3.23, width, 1.77, radius=True)
            text(slide, item["value"], x + .2, 3.6, width - .4, .65, 25, "mint" if index != 1 else "lilac", True, PP_ALIGN.CENTER)
            text(slide, item["label"], x + .2, 4.44, width - .4, .35, 16, "muted", align=PP_ALIGN.CENTER)
            if index < 2:
                connector(slide, x + width + .07, 4.09, x + width + gap - .07, 4.09, arrow=True)
        text(slide, visual["bottom"], .75, 5.36, 11.83, .48, 24, bold=True)
        text(slide, visual.get("footer", ""), .75, 6.03, 11.83, .36, 17, "muted")
    elif layout == "invitation":
        for x, key, color in [(.75, "left", "mint"), (6.81, "right", "lilac")]:
            box(slide, x, 3.34, 5.77, 1.77, radius=True)
            text(slide, visual[key], x + .25, 3.77, 5.27, .93, 28, color, True)
        text(slide, visual["bottom"], .75, 5.7, 11.83, .56, 24, bold=True)
        if visual.get("detail"):
            text(slide, visual["detail"], .75, 6.25, 11.83, .25, 12, "muted")
    elif layout == "closing":
        text(slide, visual["bottom"], .75, 3.28, 11.83, .9, 33, "mint", True)
        text(slide, visual.get("supporting_line", visual.get("supporting", visual.get("body", ""))), .75, 4.7, 11.83, .72, 25)
        text(slide, visual["detail"], .75, 5.87, 11.83, .45, 16, "muted")
    return slide


def alternative_lines(round_info, film, choice_id):
    alternatives = round_info.get("alternative_turns", {})
    if isinstance(alternatives, list):
        if choice_id == round_info["selected"]:
            return ["본문의 촬영 제안 대사를 사용한다."]
        return [f"{turn['scene']} / {turn['turn_index'] + 1}번째 발화 교체 · {production.name(turn['speaker'], film)}: {turn['text']}"
                for turn in alternatives]
    turns = alternatives.get(choice_id, [])
    if isinstance(turns, str):
        return [turns]
    return [f"{production.name(turn['speaker'], film)}: {turn['text']}" if isinstance(turn, dict) else str(turn)
            for turn in turns]


def add_notes(slide, scene, film):
    production.add_notes(slide, scene, film)
    extra = [f"{slot['id']} / {slot['label']} / {slot['detail']}" for slot in scene.get("asset_slots", [])]
    for question in film["balance"]["rounds"]:
        if question["scene"] == scene["id"]:
            extra += [f"{question['id']} / 촬영 제안 선택: {question['selected']} / 공개 발화: {question['reveal_anchor']}",
                      "선택은 실제 취향에 따라 바꿀 수 있다. 프로그램 순위나 승패가 아니다."]
            for choice in question["choices"]:
                anchor = question["reveal_anchor"] if choice["id"] == question["selected"] else question.get("alternative_reveal_anchor", "실제 발화에 맞게 확정")
                extra += [f"{choice['id']} 선택 시 공개 발화: {anchor}", *alternative_lines(question, film, choice["id"])]
    slide.notes_slide.notes_text_frame.text += "\n" + "\n".join(extra)


def make_deck(film):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = f"{film['title']} — {film['subtitle']}"
    prs.core_properties.author = "Data Center팀 교육담당 이관우"
    for scene in film["scenes"]:
        add_notes(draw_visual(prs, scene, film), scene, film)
    for appendix in film["appendix"]:
        slide = drawing.base(prs, appendix["id"], "제작 참고 / 본편 사용 제외", appendix["title"])
        text(slide, appendix["body"], .75, 3.01, 11.83, .65, 18, "muted")
        for index, item in enumerate(appendix["items"]):
            y = 3.92 + index * .59
            text(slide, item["label"], .75, y, 2.05, .42, 15, "mint", True)
            text(slide, item["value"], 2.85, y, 3.92, .43, 16, bold=True)
            text(slide, item["detail"], 6.86, y, 5.72, .44, 13, "muted")
        slide._element.set("show", "0")
        slide.notes_slide.notes_text_frame.text = json.dumps(appendix, ensure_ascii=False) + "\n" + film["source_basis"]
    return prs


def make_states(film):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "밸런스 게임 · 선택 전과 A/B 공개 화면"
    scenes = {scene["id"]: scene for scene in film["scenes"]}
    for question in film["balance"]["rounds"]:
        scene = scenes[question["scene"]]
        for state, selected in CHOICE_STATES:
            slide = draw_visual(prs, scene, film, f"{question['id']}-{state}", selected)
            add_notes(slide, scene, film)
            reveal_anchor = (question["alternative_reveal_anchor"]
                             if selected and selected != question["selected"] else question["reveal_anchor"])
            slide.notes_slide.notes_text_frame.text += (
                f"\n{question['id']} / {state} / selected={selected or 'none'}"
                + f"\n공개 장면: {question['reveal_scene']} / 이 상태의 공개 발화: {reveal_anchor}"
                + "\n기존 질문 장면 안에서 교체하는 화면. 추가 장면이나 시간이 아니다.")
    return prs


def make_documents(film):
    report = production.make_documents(film, OUT, DOCS, f"content/{PACKAGE}-film.json")
    for filename in ["01-storyboard.md", "02-script.md"]:
        path = DOCS / filename
        value = path.read_text(encoding="utf-8")
        for scene in film["scenes"]:
            heading = next(line for line in value.splitlines()
                           if line.startswith(("## ", "### ")) and scene["id"] in line)
            if scene.get("asset_slots"):
                slots = [f"- **{slot['id']} — {slot['label'].replace(chr(10), ' / ')}**: {slot['detail']}"
                         for slot in scene["asset_slots"]]
                value = value.replace(heading, heading + "\n\n**사용자가 넣을 실제 자료**\n\n" + "\n".join(slots), 1)
        value += "\n## 세 질문과 선택 화면\n\n촬영 제안 대사를 먼저 적었다. 실제 선택은 A/B 어느 쪽이든 가능하며 프로그램의 우열이나 승자를 정하지 않는다. 선택 화면 9장은 본편 시간 안에서 교체한다.\n"
        for question in film["balance"]["rounds"]:
            value += f"\n### {question['id']} · 선택 공개\n\n- 본편: {question['scene']} / 공개: {question['reveal_scene']} ‘{question['reveal_anchor']}’\n- 촬영 제안 선택: {question['selected']} / 화면: PRE → POST-A 또는 POST-B\n"
            for choice in question["choices"]:
                lines = alternative_lines(question, film, choice["id"])
                if lines:
                    anchor = question["reveal_anchor"] if choice["id"] == question["selected"] else question.get("alternative_reveal_anchor", "실제 발화에 맞게 확정")
                    value += f"\n**{choice['id']} · {choice['name']} 선택 시**\n\n공개 발화: ‘{anchor}’\n\n" + "\n\n".join(lines) + "\n"
        path.write_text(value.rstrip() + "\n", encoding="utf-8")
    return report


def main():
    film = prepare_film(json.loads(SOURCE.read_text(encoding="utf-8")))
    check_source(film)
    OUT.mkdir(parents=True, exist_ok=True)
    DOCS.mkdir(parents=True, exist_ok=True)
    report = make_documents(film)
    make_deck(film).save(OUT / f"learning-fair-2026-{PACKAGE}.pptx")
    make_states(film).save(OUT / "choice-states.pptx")
    print(json.dumps({"package": PACKAGE, "slides": 13, "choice_states": 9,
                      "spoken_syllables": report["total_spoken_syllables"],
                      "main_estimates_seconds": report["main_estimates_seconds"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
