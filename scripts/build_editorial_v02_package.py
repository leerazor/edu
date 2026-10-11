"""Build the independent editorial v02 deck with editable photo slots."""

from __future__ import annotations

import json
import re
from contextlib import contextmanager
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import build_editorial_package as shared

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "content/editorial-v02-film.json"
OUT = ROOT / "deliverables/editorial-v02"
DOCS = ROOT / "docs/editorial-v02"
DECK = OUT / "learning-fair-2026-editorial-v02.pptx"
PHOTO_INSTRUCTION = (
    "이 PPTX는 사용자가 실제 사진·화면 원본을 나중에 붙이는 제작본이다. "
    "점선 영역의 ‘자료 삽입’ 안내와 IMG ID는 교체할 편집 가능한 도형이다. "
    "빈 삽입 영역을 최종 영상으로 그대로 사용하라는 뜻이 아니다. "
    "최종 편집 전 해당 실제 자료를 삽입하고 안내 글자와 점선을 제거한다. "
    "원본을 쓰지 않으면 영상팀이 같은 메시지의 키워드 화면으로 다시 구성한다."
)


@contextmanager
def package_paths():
    """Redirect the reused document writer for this call, then restore v01."""
    destinations = {"SOURCE": SOURCE, "OUT": OUT, "DOCS": DOCS, "DECK": DECK}
    original = {key: getattr(shared, key) for key in destinations}
    try:
        for key, value in destinations.items():
            setattr(shared, key, value)
        yield
    finally:
        for key, value in original.items():
            setattr(shared, key, value)


def slot_lines(scene):
    return [f"- **{slot['id']} — {slot['label']}**: {slot.get('detail', '')}"
            for slot in scene.get("asset_slots", [])]


def add_document_slots(value, film):
    headings = list(re.finditer(r"(?m)^#{2,3}[^\n]*\b(S\d{2})\b[^\n]*$", value))
    scenes = {scene["id"]: scene for scene in film["scenes"]}
    for heading in reversed(headings):
        lines = slot_lines(scenes[heading.group(1)])
        if lines:
            insert = "\n\n**사용자가 붙일 실제 자료**\n\n" + "\n".join(lines)
            value = value[:heading.end()] + insert + value[heading.end():]
    return value


def make_documents(film):
    with package_paths():
        report = shared.make_documents(film)
    for filename in ["01-storyboard.md", "02-script.md"]:
        path = DOCS / filename
        value = path.read_text(encoding="utf-8")
        value = value.replace("# 설명형 영상 콘티", "# 공동해설 v02 콘티 · 사진 삽입용 제작본")
        value = value.replace("# 설명형 영상 촬영 대본", "# 공동해설 v02 촬영 대본 · 사진 삽입용 제작본")
        value = value.replace("editorial/rehearsal.csv", "deliverables/editorial-v02/rehearsal.csv")
        value = value.replace("content/editorial-film.json", "content/editorial-v02-film.json")
        value = value.replace("기존 기본안과 월드컵안은 별도 원본·산출물이다.",
                              "기존 기본안·월드컵안·공동해설 v01은 별도 원본·산출물이다.")
        value = value.replace(
            "현재 내부 행사 사진·화면 원본은 미수급이다. 원본이 없으면 JSON에 정의된 키워드·도식과 스튜디오 화면으로 제작한다.",
            "실제 행사 사진·화면 원본은 사용자가 별도로 삽입한다. 장면별 asset_slots의 ID·자료명·교체 방법을 확인한다.")
        value = value.replace("10월 세미나 상태는 개최 확인 후 갱신한다.",
                              "세미나 개최 상태는 이 버전 JSON과 확인된 자료를 기준으로 표시한다.")
        value = add_document_slots(value, film)
        value += "\n## 사진 삽입용 제작본 사용 방법\n\n" + PHOTO_INSTRUCTION + "\n"
        path.write_text(value, encoding="utf-8")
    return report


def draw_photo_slot(slide, page, slot, x, y, width, height):
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(width), Inches(height))
    frame.name = f"{page}-asset-frame-{slot['id']}"
    frame.fill.solid()
    frame.fill.fore_color.rgb = shared.rgb("FFFFFF")
    frame.line.color.rgb = shared.rgb("AAB8B0")
    frame.line.width = Pt(1.2)
    frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    for effect in frame._element.xpath("./p:style/a:effectRef"):
        effect.set("idx", "0")
    label = shared.draw_text(slide, "자료 삽입", x + .16, y + .4, width - .32, .36,
                             16, shared.ACCENT, True, PP_ALIGN.CENTER)
    label.name = f"{page}-asset-instruction-{slot['id']}"
    label = shared.draw_text(slide, slot["label"], x + .16, y + 1.02, width - .32,
                             1.04, 22 if width > 3 else 20,
                             shared.INK, True, PP_ALIGN.CENTER)
    label.name = f"{page}-asset-label-{slot['id']}"
    label = shared.draw_text(slide, slot["id"], x + .16, y + height - .51,
                             width - .32, .3, 14, shared.MUTED, align=PP_ALIGN.CENTER)
    label.name = f"{page}-asset-id-{slot['id']}"


def draw_asset_visual(prs, scene):
    visual = scene["visual"]
    slide = shared.base(prs, scene["slide"], visual.get("eyebrow", ""),
                        visual["headline"], visual.get("footer", ""))
    shared.draw_text(slide, visual.get("body", ""), .75, 3.45, 6.0, .88, 25)
    items = visual.get("items", [])
    assert len(items) <= 3, f"{scene['slide']}: photo layout supports at most three short fact rows"
    step = 1.86 / max(len(items), 1)
    for i, item in enumerate(items):
        y = 4.46 + i * step
        if i:
            shared.line(slide, .75, y - .065, 6.75, y - .065)
        shared.draw_text(slide, item.get("label", ""), .75, y, 1.76, .38,
                         15, shared.ACCENT, True)
        shared.draw_text(slide, item.get("value", ""), 2.66, y, 4.09, .43,
                         23, shared.INK, True)
        shared.draw_text(slide, item.get("detail", ""), 2.66, y + .43, 4.09,
                         min(.65, step - .43), 15, shared.MUTED)
    slots = scene["asset_slots"]
    gap = .16
    width = (5.3 - gap * (len(slots) - 1)) / len(slots)
    for i, slot in enumerate(slots):
        draw_photo_slot(slide, scene["slide"], slot,
                        7.25 + i * (width + gap), 3.45, width, 2.9)
    return slide


def add_notes(slide, scene, film):
    shared.add_notes(slide, scene, film)
    value = slide.notes_slide.notes_text_frame.text
    value = value.replace("내부 행사 사진·화면 원본은 현재 미수급이다.",
                          "실제 자료는 사용자가 별도로 삽입한다.")
    lines = slot_lines(scene)
    if lines:
        value += "\n자료 삽입 슬롯:\n" + "\n".join(lines) + "\n" + PHOTO_INSTRUCTION
    slide.notes_slide.notes_text_frame.text = value


def make_deck(film):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = f"{film['title']} — {film['subtitle']}"
    prs.core_properties.author = "Data Center팀"
    for scene in film["scenes"]:
        slide = (draw_asset_visual(prs, scene) if scene.get("asset_slots")
                 else shared.draw_visual(prs, scene["slide"], scene["visual"]))
        add_notes(slide, scene, film)
    for appendix in film["appendix"]:
        shared.draw_appendix(prs, appendix, film)
    return prs


def check_source(film):
    shared.check_source(film)
    for scene in film["scenes"]:
        slots = scene.get("asset_slots", [])
        assert len(slots) <= 2, f"{scene['id']}: at most two photo slots per slide"
        assert len({slot["id"] for slot in slots}) == len(slots), scene["id"]
        for slot in slots:
            assert re.fullmatch(r"IMG\d{2}", slot["id"]) and slot["label"].strip(), scene["id"]
            assert slot["id"] in scene["assets"], f"{scene['id']}: slot missing from scene assets"


def main():
    film = json.loads(SOURCE.read_text(encoding="utf-8"))
    check_source(film)
    prs = make_deck(film)
    OUT.mkdir(parents=True, exist_ok=True)
    DOCS.mkdir(parents=True, exist_ok=True)
    report = make_documents(film)
    prs.save(DECK)
    print(json.dumps({"slides": len(prs.slides), "spoken_syllables": report["total_spoken_syllables"],
                      "main_estimates_seconds": report["main_estimates_seconds"],
                      "photo_slots": sum(len(s.get("asset_slots", [])) for s in film["scenes"]),
                      "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
