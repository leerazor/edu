"""Shared navy/mint/lilac PowerPoint drawing primitives for the two packages."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

FONT = "맑은 고딕"
COLORS = {
    "bg": "101D2A", "panel": "1B2D3E", "text": "F6F8F9",
    "muted": "B4C6D4", "mint": "83E8C5", "lilac": "C1B3F4",
    "line": "385568", "amber": "F0CB7B",
}
STAGES = ["밖에서 배우다", "함께 해보다", "서로 이해하다", "다시 나누다"]


def rgb(key: str) -> RGBColor:
    return RGBColor.from_string(COLORS.get(key, key))


def text(slide, value: str, x: float, y: float, w: float, h: float,
         size: int = 24, color: str = "text", bold: bool = False,
         align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        east_asia = OxmlElement("a:ea")
        east_asia.set("typeface", FONT)
        run._r.get_or_add_rPr().append(east_asia)
    return shape


def box(slide, x, y, w, h, fill="panel", line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def connector(slide, x1, y1, x2, y2, color="mint", width=2, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        tip = OxmlElement("a:tailEnd")
        tip.set("type", "triangle")
        shape.line._get_or_add_ln().append(tip)
    return shape


def base(prs, page, label, title, active=None, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("bg")
    box(slide, .6, .53, .08, .3, "mint")
    text(slide, "2026 LEARNING FAIR  /  DATA CENTER", .86, .5, 9.7, .35, 12, "muted")
    text(slide, page, 11.64, .49, .85, .35, 13, "mint", True, PP_ALIGN.RIGHT)
    text(slide, label, .75, 1.13, 11.85, .4, 16, "mint", True)
    text(slide, title, .75, 1.67, 11.85, 1.25, 32, bold=True)
    if subtitle:
        text(slide, subtitle, .75, 2.78, 11.85, .45, 17, "muted")
    connector(slide, .75, 6.6, 12.58, 6.6, "line", 1)
    if active is not None:
        for i, stage in enumerate(STAGES):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.8 + i * 3.02), Inches(6.93), Inches(.11), Inches(.11))
            dot.fill.solid()
            dot.fill.fore_color.rgb = rgb("mint" if i <= active else "line")
            dot.line.fill.background()
            text(slide, stage, .99 + i * 3.02, 6.85, 2.65, .35, 12,
                 "text" if i == active else "muted", i == active)
    else:
        text(slide, "검토용 초안 · 회사 KV/폰트 수급 후 교체", .75, 6.86, 11.9, .35, 12, "muted")
    return slide
