"""Validate the separate editorial v04 package after actual PowerPoint export."""

import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path

import fitz
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation

from build_editorial_package import scene_content_hash
from build_editorial_v04_package import check_source
from build_package import cue_text, measure, name, timecode
from validate_worldcup_package import compact, markdown_scene_blocks, prompter_scene_blocks, prompter_turns

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/editorial-v04'
DOCS = ROOT / 'docs/editorial-v04'
DECK = OUT / 'learning-fair-2026-editorial-v04.pptx'
PDF = OUT / 'learning-fair-2026-editorial-v04.pdf'
PREVIEW = OUT / 'preview'
ERRORS, CHECKS = [], []


def check(condition, label):
    (CHECKS if condition else ERRORS).append(label)


def main():
    ERRORS.clear()
    CHECKS.clear()
    film = json.loads((ROOT / 'content/editorial-v04-film.json').read_text())
    check_source(film)
    scenes = film['scenes']
    rows = [measure(s, film) for s in scenes]
    timing = json.loads((OUT / 'timing-report.json').read_text())
    check(timing['scenes'] == rows, 'Pronunciation estimates match the current script')
    check(film['subtitle'] == '서로에게 배우는 인프라 팀', 'v04 subtitle describes the learning culture')
    check(film['schedule']['as_of'] == '2026-10-11', 'Submission planning date is October 11')
    expected_dates = {'outsystems':'2026-10-12', 'submission':'2026-10-14', 'tech_x_final':'2026-10-16', 'filming':'2026-10-19'}
    for key, date in expected_dates.items():
        check(film['schedule'][key] == {'date':date, 'status':'scheduled'}, f'{key} confirmed schedule, pending event')
    intro = ' '.join(t['text'] for t in scenes[0]['turns'])
    check(all(term in intro for term in ['DS부문', '모든', '인프라']),
          'S01 states the team role across all DS infrastructure')
    check(all(term in intro for term in ['AI', '운영', '효율']),
          'S01 states AI use for infrastructure operation efficiency')
    check(not re.search(r'160|백육십', intro + scenes[0]['visual']['headline'] + scenes[0]['visual'].get('body', '')),
          'S01 introduction omits the 160-person size')
    check(scenes[0]['visual']['layout'] == 'process', 'P01 introduces the learning culture as a process')
    culture = intro + json.dumps(scenes[0]['visual'], ensure_ascii=False)
    check('교류' in intro and '성장' in intro, 'S01 states growth through exchange')
    check(all(any(term in culture for term in family) for family in
              [['이해', '알고', '알아', '배우', '배움'], ['참여', '함께', '경험'], ['공유', '나누', '나눔']]),
          'S01 introduces learning, shared participation, and sharing experience')
    check({'IMG02', 'IMG03'}.issubset(scenes[1]['assets']), 'S02 opens the examples with AI 4 AI seminars')
    check({'IMG04', 'IMG05'}.issubset(scenes[2]['assets']), 'S03 presents Job Fair')
    check('IMG06' in scenes[3]['assets'], 'S04 presents DCIS')
    check('IMG07' in scenes[4]['assets'], 'S05 presents Tech X recruitment')
    check(scenes[5]['assets'] == ['IMG15'], 'S06 personal GPU project evidence uses IMG15')
    check('다른 과제' in ' '.join(t['text'] for t in scenes[5]['turns']), 'Personal GPU project is distinguished from IAP')
    recruitment = ' '.join(t['text'] for t in scenes[4]['turns'])
    check('3월' in recruitment and '4월' in recruitment, 'S05 recruitment and project start both specified')
    check(timing['iap_allocated_seconds'] == 35, 'IAP is one 35-second example')
    iap_scenes = [s['id'] for s in scenes if 'IAP' in ' '.join(t['text'] for t in s['turns'])]
    check(iap_scenes == ['S07'] and film.get('iap_scene') == 'S07', 'IAP is spoken only in S07')
    check(scenes[7]['assets'] == ['IMG13'], 'Platform evidence uses IMG13, not cooling assets')
    check(sum(len(s.get('asset_slots', [])) for s in scenes) == 9, 'Exactly nine editable material slots')
    for rate in film['reading_rates']:
        total = round(sum(r['estimates_seconds'][rate] for r in rows), 2)
        check(timing['main_estimates_seconds'][rate] == total, f'{rate} estimate total')
        check(timing['whole_film_estimates_seconds'][rate] == round(total + 25, 2), f'{rate} fixed segments')
    check(timing['whole_film_estimates_seconds']['slow'] < 600, 'Usable take estimate below ten minutes')

    script = (DOCS / '02-script.md').read_text()
    board = (DOCS / '01-storyboard.md').read_text()
    prompt = (OUT / 'teleprompter.txt').read_text()
    script_blocks, board_blocks = markdown_scene_blocks(script), markdown_scene_blocks(board)
    prompt_blocks = prompter_scene_blocks(prompt)
    registered = set(re.findall(r'(?m)^\|\s*(IMG\d{2})\s*\|', (ROOT / 'docs/03-facts-and-assets.md').read_text()))
    for doc in DOCS.glob('*.md'):
        check('이 문서에서 해야 할 일' in doc.read_text(), f'{doc.name} reader action')
    with (OUT / 'rehearsal.csv').open(encoding='utf-8-sig', newline='') as handle:
        rehearsal = list(csv.DictReader(handle))
    check([r['scene'] for r in rehearsal] == [s['id'] for s in scenes], 'Rehearsal scene order')
    for s, r in zip(scenes, rehearsal, strict=True):
        sid = s['id']
        check(set(re.findall(r'P\d{2}', s['cue']['effect'])) == {s['slide']},
              f'{sid} main editing effect points to its current slide after reordering')
        if sid == 'S06':
            check('S07의 IAP' in s['hold'] and 'S06의 IAP' not in s['hold'],
                  'GPU experience handoff points to IAP in S07')
        check(set(s['assets']).issubset(registered), f'{sid} asset IDs')
        check(s['camera'] in board_blocks[sid] and s['hold'] in board_blocks[sid], f'{sid} storyboard direction')
        check(s['cue']['anchor'] in board_blocks[sid], f'{sid} storyboard cue')
        for turn in s['turns']:
            expected = f"{name(turn['speaker'], film)}: {cue_text(turn, s, film)}"
            check(compact(expected) in compact(script_blocks[sid].replace('**', '')), f'{sid} script synchronization')
        actual_slide, block = prompt_blocks[sid]
        check(actual_slide == s['slide'], f'{sid} prompter slide')
        expected = [(name(t['speaker'], film), compact(cue_text(t, s, film, phonetic=True))) for t in s['turns']]
        check([(who, compact(text)) for who, text in prompter_turns(block)] == expected, f'{sid} prompter pronunciation and cues')
        check(r['content_hash'] == scene_content_hash(s, film), f'{sid} rehearsal content hash')
        check(r['start'] == timecode(s['start']) and r['end'] == timecode(s['start'] + s['duration']), f'{sid} rehearsal times')

    with zipfile.ZipFile(DECK) as archive:
        check(archive.testzip() is None, 'PPTX archive integrity')
        check(not any(b'<p:timing' in archive.read(p) for p in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', p)), 'Static deck: animation instructions only')
    deck = Presentation(DECK)
    ids = [s['slide'] for s in scenes] + [a['id'] for a in film['appendix']]
    check(len(deck.slides) == 13, '10 main slides and 3 appendices')
    check(abs(deck.slide_width / deck.slide_height - 16 / 9) < .001, 'PPTX aspect ratio')
    for i, slide in enumerate(deck.slides):
        pid = ids[i]
        check(slide._element.get('show', '1') == ('1' if i < 10 else '0'), f'{pid} visibility')
        notes = slide.notes_slide.notes_text_frame.text
        check(bool(notes), f'{pid} speaker notes')
        if i < 10:
            s = scenes[i]
            check(s['camera'] in notes and s['hold'] in notes, f'{pid} notes directions')
            for t in s['turns']:
                check(cue_text(t, s, film) in notes, f'{pid} notes dialogue and cue')
            for cue in s.get('secondary_cues', []):
                check(all(cue[k] in notes for k in ['anchor', 'label', 'effect']), f'{pid} secondary cue')
            for slot in s.get('asset_slots', []):
                shapes = {shape.name: shape for shape in slide.shapes}
                check(f"{pid}-asset-frame-{slot['id']}" in shapes, f"{pid} editable slot frame {slot['id']}")
                label = shapes.get(f"{pid}-asset-label-{slot['id']}")
                check(label is not None and label.text == slot['label'], f"{pid} visible material label {slot['id']}")
                check(all(slot[k] in notes for k in ['id','label','detail']), f"{pid} material instructions in notes")
                check(slot['label'] in script_blocks[s['id']] and slot['detail'] in board_blocks[s['id']], f"{pid} material instructions in documents")
        for shape in slide.shapes:
            check(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= deck.slide_width + 20 and shape.top + shape.height <= deck.slide_height + 20, f'{pid} shape bounds')
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            check(run.font.name == '맑은 고딕' and run.font.size is not None, f'{pid} explicit font')

    expected_pngs = {f'slide-{i:02}.png' for i in range(1, 14)}
    check({p.name for p in PREVIEW.glob('slide-*.png')} == expected_pngs, 'Exactly 13 actual PNG exports')
    with fitz.open(PDF) as pdf:
        check(len(pdf) == len(deck.slides), 'Actual PDF page count')
        for i, page in enumerate(pdf):
            extracted = compact(page.get_text())
            check(ids[i] in extracted and '�' not in extracted, f'{ids[i]} PDF text')
            for shape in deck.slides[i].shapes:
                if shape.has_text_frame and shape.text.strip():
                    check(compact(shape.text) in extracted, f'{ids[i]} all slide text rendered: {shape.text[:20]}')
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        x0,y0,x1,y1=span['bbox']
                        check(x0 >= -1 and y0 >= -1 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1, f'{ids[i]} PDF text bounds')
            with Image.open(PREVIEW / f'slide-{i+1:02}.png') as image:
                check(image.size == (1600,900), f'{ids[i]} PNG resolution')
                pix = page.get_pixmap(matrix=fitz.Matrix(.2,.2), alpha=False)
                ref = Image.frombytes('RGB', (pix.width,pix.height), pix.samples)
                a=image.convert('L').resize((96,54),Image.Resampling.LANCZOS)
                b=ref.convert('L').resize((96,54),Image.Resampling.LANCZOS)
                check(ImageStat.Stat(ImageChops.difference(a,b)).mean[0]/255 < .12, f'{ids[i]} PDF and PNG agree')
    baseline = json.loads((OUT / 'preservation-baseline.json').read_text())
    for path, digest in baseline['files'].items():
        check(hashlib.sha256((ROOT/path).read_bytes()).hexdigest() == digest, f'Preserved: {path}')
    canvas = Image.new('RGB', (1640, 1730), '#E8EBE5')
    draw = ImageDraw.Draw(canvas)
    for i in range(13):
        x,y=20+(i%3)*540,20+(i//3)*340
        with Image.open(PREVIEW/f'slide-{i+1:02}.png') as image:
            canvas.paste(image.convert('RGB').resize((520,293),Image.Resampling.LANCZOS),(x,y))
        draw.text((x,y+303), ids[i] + (' / MAIN' if i<10 else ' / HIDDEN APPENDIX'), fill='#202522')
    canvas.save(PREVIEW/'contact-sheet.jpg',quality=93)
    artifacts = [ROOT/'content/editorial-v04-film.json', DECK, PDF, DOCS/'01-storyboard.md', DOCS/'02-script.md', OUT/'teleprompter.txt', OUT/'timing-report.json']
    report={'status':'FAIL' if ERRORS else 'PASS','passed_checks':len(CHECKS),'errors':ERRORS,
            'metrics':{'main_seconds':275,'whole_seconds':300,'iap_seconds':35,'spoken_syllables':timing['total_spoken_syllables'],'estimates':timing['whole_film_estimates_seconds'],'measured_scenes':sum(bool(r['measured_seconds']) for r in rehearsal),'preserved_files':len(baseline['files']),'asset_slots':sum(len(s.get('asset_slots',[])) for s in scenes)},
            'artifact_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in artifacts},
            'limits':['Reading estimates are not presenter rehearsal measurements.','Internal event images and company KV are not supplied.','Editorial review is not a viewer study. Motion remains editing instructions.']}
    (OUT/'validation-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False))
    if ERRORS:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
