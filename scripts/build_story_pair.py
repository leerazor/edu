"""Build two independent revisions using the original deck's visual language."""

from __future__ import annotations

import copy
import json
from contextlib import contextmanager
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import build_package as base_lib
import build_editorial_package as document_lib
from build_package import box, connector, measure, text

ROOT = Path(__file__).resolve().parents[1]
VARIANTS = {"story": "story-v02", "worldcup": "worldcup-v02"}


def check_source(film):
    document_lib.check_source(film)
    assert film["official_expansion"] == "Advanced Infrastructure for Artificial Intelligence"
    assert film["variant"] in VARIANTS
    iap = [s for s in film["scenes"] if "IAP" in " ".join(t["text"] for t in s["turns"])]
    assert len(iap) == 1 and iap[0]["id"] == film["iap_scene"]
    assert iap[0]["duration"] <= 20
    if film["variant"] == "worldcup":
        tournament = film["tournament"]
        assert tournament["actual_vote"] is False
        candidates = {c["id"] for c in tournament["candidates"]}
        assert len(candidates) == 4
        for match in tournament["matches"]:
            assert set(match["candidates"]) <= candidates
            assert match["winner"] in match["candidates"]
            target = next(s for s in film["scenes"] if s["id"] == match["reveal_scene"])
            raw = " ".join(t["text"] for t in target["turns"])
            assert raw.count(match["reveal_anchor"]) == 1
        assert set(tournament["matches"][2]["candidates"]) == {
            tournament["matches"][0]["winner"], tournament["matches"][1]["winner"]}


@contextmanager
def document_paths(key):
    destinations = {"SOURCE": ROOT / "content" / f"{key}-film.json",
                    "OUT": ROOT / "deliverables" / key,
                    "DOCS": ROOT / "docs" / key,
                    "DECK": ROOT / "deliverables" / key / f"learning-fair-2026-{key}.pptx"}
    saved = {k: getattr(document_lib, k) for k in destinations}
    try:
        for k, value in destinations.items():
            setattr(document_lib, k, value)
        yield
    finally:
        for k, value in saved.items():
            setattr(document_lib, k, value)


def make_documents(film, key):
    with document_paths(key):
        report = document_lib.make_documents(film)
    report["iap_allocated_seconds"] = next(s["duration"] for s in film["scenes"]
                                           if s["id"] == film["iap_scene"])
    out, docs = ROOT / "deliverables" / key, ROOT / "docs" / key
    (out / "timing-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    label = "흐름형" if film["variant"] == "story" else "월드컵형"
    for filename in ["01-storyboard.md", "02-script.md"]:
        p = docs / filename
        value = p.read_text()
        value = value.replace("설명형 영상", f"초안 기반 {label} 영상")
        value = value.replace("editorial/rehearsal.csv", f"deliverables/{key}/rehearsal.csv")
        value = value.replace("content/editorial-film.json", f"content/{key}-film.json")
        value = value.replace("IAP는 대표 사례로 35~40초를 유지하고, 초과 시 다른 장면의 중복 설명부터 조정한다.",
                              "IAP는 20초 이내의 짧은 대표 사례다. 초과하면 기능 해설을 늘리지 않고 중복 문구를 줄인다.")
        value = value.replace("기존 기본안과 월드컵안은 별도 원본·산출물이다.",
                              "이전 기본본·월드컵본·공동 해설 v01–v05는 별도 원본·산출물로 보존한다.")
        for s in film["scenes"]:
            heading = next((line for line in value.splitlines()
                            if line.startswith(("## ", "### ")) and s["id"] in line), None)
            if heading and s.get("asset_slots"):
                lines = [f"- **{slot['id']} — {slot['label'].replace(chr(10), ' / ')}**: {slot['detail']}"
                         for slot in s["asset_slots"]]
                value = value.replace(heading, heading + "\n\n**사용자가 넣을 실제 자료**\n\n" + "\n".join(lines), 1)
        if film["variant"] == "worldcup":
            rule = film["tournament"]
            value += ("\n## 월드컵 선택과 공개 화면\n\n"
                      + f"**세 경기의 공통 기준:** {rule['criterion']}\n\n"
                      + rule["disclaimer"] + "\n\n"
                      + "choice-states.pptx는 M01·M02·M03의 선택 전·후 화면 6장이다. 본편에 추가할 장면이나 시간이 아니다. 대본의 선택 공개 큐에서 해당 상태로 교체한다.\n")
        p.write_text(value)
    return report


def draw_slot(slide, pid, slot, x, y, w, h):
    frame = box(slide, x, y, w, h, line="line", radius=True)
    frame.name = f"{pid}-asset-frame-{slot['id']}"
    compact = h < 1.6
    instruction = text(slide, "실제 자료 삽입", x + .2, y + .12, w - .4, .25, 12, "mint", True)
    instruction.name = f"{pid}-asset-instruction-{slot['id']}"
    label = text(slide, slot["label"], x + .2, y + (.4 if compact else .72), w - .4,
                 .65 if compact else .98, 16 if compact else 22, bold=True)
    label.name = f"{pid}-asset-label-{slot['id']}"
    marker = text(slide, slot["id"], x + .2, y + h - .29, w - .4, .22, 10, "muted")
    marker.name = f"{pid}-asset-id-{slot['id']}"


def cards(slide, items, y=3.38, h=1.35):
    width = (11.83 - .42 * (len(items) - 1)) / len(items)
    for i, item in enumerate(items):
        x = .75 + i * (width + .42)
        box(slide, x, y, width, h, radius=True)
        text(slide, item["value"], x + .2, y + .25, width - .4, .62,
             24 if len(items) < 4 else 27, "lilac" if i == 1 else "mint", True,
             PP_ALIGN.CENTER if len(items) == 4 else PP_ALIGN.LEFT)
        text(slide, item.get("label", ""), x + .2, y + h - .48, width - .4, .34, 15, "muted")
        if i < len(items) - 1:
            connector(slide, x + width + .07, y + h / 2, x + width + .35, y + h / 2, arrow=True)


def draw_visual(prs, scene, film, pid=None, selected=None):
    v, pid = scene["visual"], pid or scene["slide"]
    previous_stages = base_lib.STAGES
    try:
        base_lib.STAGES = film["stages"]
        slide = base_lib.base(prs, pid, v["eyebrow"], v["headline"], v.get("stage"))
    finally:
        base_lib.STAGES = previous_stages
    layout = v["layout"]
    slots = scene.get("asset_slots", [])
    if pid.startswith("M"):
        marker = next(s for s in slide.shapes if s.has_text_frame and s.text == pid)
        marker.left, marker.width = Inches(11.2), Inches(1.38)
        for paragraph in marker.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(12)
    if layout == "scale":
        for i, item in enumerate(v["items"]):
            x = .75 + i * 4.02
            box(slide, x, 3.35, 3.79, 1.64, radius=True)
            text(slide, item["value"], x + .23, 3.73, 3.33, .6,
                 25 if len(item["value"]) > 8 else 29, "mint", True)
            text(slide, item["label"], x + .23, 4.49, 3.33, .36, 16, "muted")
        text(slide, v["bottom"], .75, 5.45, 11.83, .65, 26, bold=True)
    elif layout == "relationship":
        for x, value, color in [(.75, v["left"], "mint"), (7.21, v["right"], "lilac")]:
            box(slide, x, 3.27, 5.37, 1.36, radius=True)
            text(slide, value, x + .25, 3.68, 4.87, .66, 25, color, True)
        connector(slide, 6.25, 3.65, 7.04, 3.65, arrow=True)
        connector(slide, 7.04, 4.16, 6.25, 4.16, "lilac", arrow=True)
        text(slide, v["bottom"], .75, 4.95, 11.83, 1.0, 24, bold=True)
        text(slide, film["official_expansion"], .75, 6.13, 11.83, .3, 13, "muted")
    elif layout == "evidence":
        text(slide, v.get("body", ""), .75, 3.3, 5.3, .9, 24, "mint", True)
        for i, item in enumerate(v.get("items", [])):
            y = 4.35 + i * .88
            text(slide, item["value"], .75, y, 5.3, .43, 22, bold=True)
            text(slide, item.get("label", ""), .75, y + .44, 5.3, .3, 15, "muted")
        if v.get("bottom"):
            text(slide, v["bottom"], .75, 5.27, 5.3, .9, 22, bold=True)
        draw_slot(slide, pid, slots[0], 6.45, 3.28, 6.13, 2.92)
        if v.get("footer"):
            text(slide, v["footer"], .75, 6.2, 5.3, .28, 12, "muted")
    elif layout == "recruit":
        cards(slide, v["items"], 3.3)
        text(slide, "본래 부서 밖의 관심 과제에 참여", .75, 4.98, 6.18, .6, 21, bold=True)
        text(slide, v["bottom"], .75, 5.69, 6.18, .64, 16, "muted")
        draw_slot(slide, pid, slots[0], 7.25, 4.88, 5.33, 1.4)
    elif layout == "photos":
        text(slide, v["body"], .75, 2.96, 11.83, .43, 22, "mint", True)
        for i, slot in enumerate(slots):
            draw_slot(slide, pid, slot, .75 + 6.05 * i, 3.57, 5.78, 2.14)
        text(slide, v["bottom"], .75, 6.03, 11.83, .4, 19, "muted")
    elif layout == "process":
        cards(slide, v["items"], 3.66, 1.62)
        text(slide, v["bottom"], .75, 5.69, 11.83, .5, 22, bold=True)
    elif layout == "cycle":
        cards(slide, v["items"], 3.55, 1.35)
        connector(slide, 11.2, 5.05, 11.2, 5.3, "lilac")
        connector(slide, 11.2, 5.3, 2.05, 5.3, "lilac")
        connector(slide, 2.05, 5.3, 2.05, 5.05, "lilac", arrow=True)
        text(slide, v["bottom"], .75, 5.76, 11.83, .6, 24, bold=True)
    elif layout == "closing":
        title = next(s for s in slide.shapes if s.has_text_frame and s.text == v["headline"])
        for paragraph in title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(58)
        text(slide, v["bottom"], .75, 3.49, 11.83, .97, 35, "mint", True)
        text(slide, "밖에서 배우고 · 함께 해보고 · 서로 이해하고 · 다시 나눕니다", .75, 5.16, 11.83, .6, 22)
        text(slide, v["detail"], .75, 6.0, 11.83, .35, 16, "muted")
    elif layout == "duel":
        text(slide, film["tournament"]["criterion"], .75, 2.98, 11.83, .43, 18, "muted")
        match = next(m for m in film["tournament"]["matches"] if m["id"] == v["match"])
        for i, item in enumerate(v["items"]):
            x = .75 + i * 6.06
            box(slide, x, 3.63, 5.77, 2.23, radius=True,
                line="mint" if selected == match["candidates"][i] else "line")
            text(slide, item["label"], x + .25, 3.88, 5.27, .45, 24,
                 "mint" if i == 0 else "lilac", True)
            text(slide, item["value"], x + .25, 4.61, 5.27, .66, 23, bold=True)
            text(slide, item["detail"], x + .25, 5.37, 5.27, .31, 14, "muted")
        if selected:
            names = {c["id"]: c["name"] for c in film["tournament"]["candidates"]}
            text(slide, f"오늘의 추천 · {names[selected]}", .75, 6.11, 11.83, .34, 20, "mint", True)
        else:
            text(slide, "두 사람의 선택 이유를 듣고 추천합니다", .75, 6.11, 11.83, .34, 16, "muted")
    elif layout == "bracket":
        text(slide, v["criterion"], .75, 3.02, 11.83, .55, 22, "mint", True)
        names = {c["id"]: c["name"] for c in film["tournament"]["candidates"]}
        for i, match in enumerate(film["tournament"]["matches"][:2]):
            y = 3.83 + i * 1.15
            text(slide, match["round"], .75, y + .16, 1.6, .4, 17, "muted")
            for j, candidate in enumerate(match["candidates"]):
                x = 2.72 + j * 5.05
                box(slide, x, y, 4.83, .83, radius=True)
                text(slide, names[candidate], x + .2, y + .21, 4.43, .45, 23,
                     "mint" if j == 0 else "lilac", True)
        text(slide, v["bottom"], .75, 6.22, 11.83, .3, 16, "muted")
    else:
        raise ValueError(layout)
    return slide


def add_notes(slide, scene, film):
    document_lib.add_notes(slide, scene, film)
    value = slide.notes_slide.notes_text_frame.text
    value += "\n" + "\n".join(f"{s['id']} / {s['label']} / {s['detail']}" for s in scene.get("asset_slots", []))
    if film["variant"] == "worldcup":
        value += "\n" + film["tournament"]["criterion"] + "\n" + film["tournament"]["disclaimer"]
    value += "\n" + "\n".join(f"{s['title']}: {s['url']}" for s in film["external_sources"])
    slide.notes_slide.notes_text_frame.text = value


def make_deck(film):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = f"{film['title']} — {film['subtitle']}"
    prs.core_properties.author = "Data Center팀 교육담당 이관우"
    for scene in film["scenes"]:
        slide = draw_visual(prs, scene, film)
        add_notes(slide, scene, film)
    for appendix in film["appendix"]:
        slide = base_lib.base(prs, appendix["id"], "제작 참고 / 본편 사용 제외", appendix["title"])
        text(slide, appendix["body"], .75, 3.01, 11.83, .65, 18, "muted")
        for i, item in enumerate(appendix["items"]):
            y = 3.93 + i * .57
            text(slide, item["label"], .75, y, 2.05, .38, 16, "mint", True)
            text(slide, item["value"], 2.85, y, 3.92, .43, 17, bold=True)
            text(slide, item["detail"], 6.86, y, 5.72, .44, 14, "muted")
        slide._element.set("show", "0")
        slide.notes_slide.notes_text_frame.text = json.dumps(appendix, ensure_ascii=False) + "\n" + film["source_basis"]
        if appendix["id"] == "A02":
            slide.notes_slide.notes_text_frame.text += "\n" + json.dumps(film["external_sources"], ensure_ascii=False)
            if film["variant"] == "worldcup":
                slide.notes_slide.notes_text_frame.text += "\n" + film["tournament"]["disclaimer"]
    return prs


def make_states(film, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "성장 월드컵 · 경기별 선택 전후 화면"
    scenes = {s["id"]: s for s in film["scenes"]}
    for match in film["tournament"]["matches"]:
        for after in [False, True]:
            state = "선택 후" if after else "선택 전"
            slide = draw_visual(prs, scenes[match["scene"]], film,
                                f"{match['id']}-{'POST' if after else 'PRE'}",
                                match["winner"] if after else None)
            slide.notes_slide.notes_text_frame.text = (
                f"{match['id']} / {state} / winner={match['winner']}\n"
                + film["tournament"]["criterion"] + "\n" + film["tournament"]["disclaimer"]
                + f"\n{match['reveal_scene']} / 공개 발화: {match['reveal_anchor']}\n"
                + "본편 배정 시간 안에서 교체하는 화면. 추가 장면이나 시간이 아님.")
    prs.save(out / "choice-states.pptx")


def main():
    for variant, key in VARIANTS.items():
        film = json.loads((ROOT / "content" / f"{key}-film.json").read_text())
        check_source(film)
        out = ROOT / "deliverables" / key
        out.mkdir(parents=True, exist_ok=True)
        (ROOT / "docs" / key).mkdir(parents=True, exist_ok=True)
        report = make_documents(film, key)
        make_deck(film).save(out / f"learning-fair-2026-{key}.pptx")
        if variant == "worldcup":
            make_states(film, out)
        print(json.dumps({"variant": variant, "slides": 13,
                          "spoken_syllables": report["total_spoken_syllables"],
                          "main_estimates_seconds": report["main_estimates_seconds"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
