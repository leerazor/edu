"""Build the separate editorial film package from its JSON source."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_package import (
    connector as common_connector,
    cue_text,
    measure,
    name,
    rgb,
    text as common_text,
    timecode,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "content/editorial-film.json"
OUT = ROOT / "deliverables/editorial"
DOCS = ROOT / "docs/editorial"
DECK = OUT / "learning-fair-2026-editorial.pptx"
BG, INK, ACCENT, MUTED, RULE = "FAFAF7", "202522", "176B5B", "626962", "D9DDD7"
LAYOUTS = {"title", "statement", "facts", "process", "split"}


def draw_text(slide, value, x, y, w, h, size=24, color=INK, bold=False,
              align=PP_ALIGN.LEFT):
    return common_text(slide, str(value), x, y, w, h, size, color, bold, align)


def line(slide, x1, y1, x2, y2, color=RULE, width=1, arrow=False):
    shape = common_connector(slide, x1, y1, x2, y2, color, width, arrow)
    for effect in shape._element.xpath('./p:style/a:effectRef'):
        effect.set('idx', '0')
    return shape


def base(prs, page, eyebrow, headline, footer="", title=False, statement=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    draw_text(slide, eyebrow, .75, .62, 11.8, .48, 18, ACCENT, True)
    draw_text(slide, headline, .75, 2.0 if statement else 1.43, 11.8, 1.9,
              48 if statement else 44 if title else 40, INK, True)
    if footer:
        draw_text(slide, footer, .75, 6.63, 11.8, .33, 12, MUTED)
    line(slide, .75, 7.02, 12.58, 7.02)
    draw_text(slide, page, .75, 7.14, .85, .23, 10, ACCENT, True)
    draw_text(slide, "2026 LEARNING FAIR", 8.2, 7.14, 4.38, .23,
              10, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def draw_rows(slide, items, top=4.68, height=1.82):
    """Short editorial rows; all words come from the JSON visual object."""
    count = len(items)
    if not count:
        return
    step = height / count
    for i, item in enumerate(items):
        y = top + i * step
        if i:
            line(slide, .75, y - .08, 12.58, y - .08)
        draw_text(slide, item.get("label", ""), .75, y, 2.35, step - .08,
                  18, ACCENT, True)
        draw_text(slide, item.get("value", ""), 3.23, y, 4.25, step - .08,
                  23, INK, True)
        draw_text(slide, item.get("detail", ""), 7.67, y, 4.91, step - .08,
                  18, MUTED)


def draw_columns(slide, items, layout):
    count = len(items)
    if not count:
        return
    gap = .45
    width = (11.83 - gap * (count - 1)) / count
    for i, item in enumerate(items):
        x = .75 + i * (width + gap)
        if i:
            line(slide, x - gap / 2, 4.24, x - gap / 2, 6.48)
        draw_text(slide, item.get("label", ""), x, 4.26, width, .5,
                  18, ACCENT, True)
        value_size = 54 if layout == "facts" else 30 if layout == "split" else 27
        draw_text(slide, item.get("value", ""), x, 4.87, width, 1.04,
                  value_size, INK, True)
        draw_text(slide, item.get("detail", ""), x, 5.99, width, .54,
                  20 if count < 4 else 18, MUTED)


def draw_visual(prs, page, visual):
    layout = visual["layout"]
    assert layout in LAYOUTS, f"{page}: unsupported layout {layout}"
    items = visual.get("items", [])
    limit = 2 if layout == "split" else 3 if layout in {"title", "statement"} else 4
    assert len(items) <= limit, f"{page}: {layout} supports at most {limit} items"
    assert str(visual["headline"]).count("\n") <= 1, f"{page}: headline exceeds two explicit lines"
    statement = layout == "statement" and not items
    slide = base(prs, page, visual.get("eyebrow", ""), visual["headline"],
                 visual.get("footer", ""), title=layout == "title", statement=statement)
    body = visual.get("body", "")
    if layout == "title":
        draw_text(slide, body, .75, 3.65, 11.8, .94, 29)
        draw_rows(slide, items, top=4.93, height=1.55)
    elif layout == "statement":
        draw_text(slide, body, .75, 4.76 if statement else 3.46, 11.8, 1.02, 28)
        draw_rows(slide, items)
    elif layout == "process" and visual.get("flow"):
        draw_text(slide, body, .75, 3.35, 11.8, .7, 25)
        for i, item in enumerate(items):
            x = .75 + i * 4.1
            draw_text(slide, item.get("label", ""), x, 4.02, 3.55, .5, 18, ACCENT, True)
            draw_text(slide, item.get("value", ""), x, 4.68, 3.55, .8, 34, INK, True)
            draw_text(slide, item.get("detail", ""), x, 5.68, 3.55, .65, 20, MUTED)
            if i + 1 < len(items):
                line(slide, x + 3.32, 5.03, x + 3.88, 5.03, ACCENT, 1.5, arrow=True)
    elif layout == "facts" and len(items) == 1:
        draw_text(slide, body, .75, 3.35, 6.9, 1.1, 27)
        item = items[0]
        line(slide, 8.22, 3.55, 8.22, 6.2)
        draw_text(slide, item.get("label", ""), 8.7, 3.57, 3.88, .5, 18, ACCENT, True)
        draw_text(slide, item.get("value", ""), 8.63, 4.16, 3.95, 1.6, 78, INK, True)
        draw_text(slide, item.get("detail", ""), .75, 5.45, 6.9, .9, 24, MUTED)
    else:
        draw_text(slide, body, .75, 3.35, 11.8, .7, 25)
        draw_columns(slide, items, layout)
    return slide


def add_notes(slide, scene, film):
    lines = [
        f"{scene['slide']} / {scene['id']} / 최종 {timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}",
        f"원본 버전: {film['version']} / {film['date']}",
        "자료 큐는 발화 기준. 타임코드는 최종 편집 목표이며 실측 발화 시점이 아니다.",
        f"자료 등장: {scene['cue']['anchor']}", scene["cue"]["effect"],
        f"촬영: {scene['camera']}",
        f"원본: {', '.join(scene['assets']) or '편집 가능한 키워드·도식'}",
        "스크립트:",
    ]
    lines += [f"{name(turn['speaker'], film)}: {cue_text(turn, scene, film)}"
              for turn in scene["turns"]]
    lines += [f"추가 큐: {cue['anchor']} / {cue['label']} / {cue['effect']}"
              for cue in scene.get("secondary_cues", [])]
    lines += [scene["hold"], film.get("source_basis", ""), "PPTX에는 실제 애니메이션이 없다. 효과와 자료 전환은 영상팀 편집 지시다.",
              "내부 행사 사진·화면 원본은 현재 미수급이다. 범용 외부 사진으로 실제 행사를 대체하지 않는다.",
              "출처: 사용자 브리프·추가 확인. 사실과 원본 ID 기준: docs/03-facts-and-assets.md."]
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def draw_appendix(prs, appendix, film):
    slide = base(prs, appendix["id"], "제작 참고 / 본편 사용 제외", appendix["title"])
    draw_text(slide, appendix.get("body", ""), .75, 3.27, 11.8, .72, 19)
    items = appendix.get("items", [])
    assert len(items) <= 8, f"{appendix['id']}: appendix supports at most eight rows"
    step = 2.5 / max(len(items), 1)
    for i, item in enumerate(items):
        y = 4.14 + i * step
        if i:
            line(slide, .75, y - .055, 12.58, y - .055)
        draw_text(slide, item.get("label", ""), .75, y, 2.04, step - .065,
                  15 if len(items) < 6 else 13, ACCENT, True)
        draw_text(slide, item.get("value", ""), 2.96, y, 3.54, step - .065,
                  17 if len(items) < 6 else 14, INK, True)
        draw_text(slide, item.get("detail", ""), 6.69, y, 5.89, step - .065,
                  16 if len(items) < 6 else 13, MUTED)
    slide._element.set("show", "0")
    slide.notes_slide.notes_text_frame.text = "\n".join([
        f"{appendix['id']} / 숨김 제작 부록 / {film['date']}",
        appendix["title"], appendix.get("body", ""), appendix.get("note", ""),
        *[" / ".join(str(item.get(key, "")) for key in ["label", "value", "detail"])
          for item in items],
        "본편 사용 제외. 검토 PDF·PNG에는 이 숨김 부록도 포함된다.",
    ])
    return slide


def make_deck(film):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = f"{film['title']} — {film['subtitle']}"
    prs.core_properties.author = "Data Center팀"
    for scene in film["scenes"]:
        slide = draw_visual(prs, scene["slide"], scene["visual"])
        add_notes(slide, scene, film)
    for appendix in film["appendix"]:
        draw_appendix(prs, appendix, film)
    return prs


def scene_content_hash(scene, film):
    """A measurement belongs to this speech and cue configuration only."""
    payload = {key: scene.get(key) for key in
               ["id", "slide", "turns", "cue", "secondary_cues", "duration", "pause_seconds"]}
    payload["pronunciations"] = film["pronunciations"]
    payload["presenters"] = film["presenters"]
    value = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def cell(value):
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def make_documents(film):
    rows = [measure(scene, film) for scene in film["scenes"]]
    report = {
        "basis": "Planning estimate, not a rehearsal measurement",
        "version": film["version"], "date": film["date"],
        "reading_rates_syllables_per_second": film["reading_rates"],
        "main_allocated_seconds": sum(s["duration"] for s in film["scenes"]),
        "fixed_seconds": sum(s["duration"] for s in film["fixed_segments"]),
        "scenes": rows,
        "total_spoken_syllables": sum(r["spoken_syllables"] for r in rows),
        "main_estimates_seconds": {key: round(sum(r["estimates_seconds"][key] for r in rows), 2)
                                   for key in film["reading_rates"]},
    }
    report["whole_film_estimates_seconds"] = {
        key: round(value + report["fixed_seconds"], 2)
        for key, value in report["main_estimates_seconds"].items()
    }
    report["iap_allocated_seconds"] = sum(
        scene["duration"] for scene in film["scenes"]
        if re.search(r"(?<![A-Za-z])IAP(?![A-Za-z])|Infra Agent Platform",
                     " ".join(turn["text"] for turn in scene["turns"]))
    )
    (OUT / "timing-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    heading = f"**{film['version']} · {film['date']}** · {film['title']} — {film['subtitle']}"
    board = ["# 설명형 영상 콘티", "",
             "> 이 문서에서 해야 할 일: 장면별 자료 등장 발화·화면·촬영·홀드를 확인하고, 같은 발화 큐에 맞춰 편집한다.", "", heading, "",
             "전체 300초 = 회사 인트로 5초 + AI 팀 소개 15초 + 스튜디오 본편 275초 + 회사 아웃트로 5초.", "",
             "| 최종 타임코드 | 길이 | 장면 | 자료 | 화면 목적 |", "| --- | ---: | --- | --- | --- |"]
    timeline = sorted(film["fixed_segments"] + film["scenes"], key=lambda scene: scene["start"])
    for scene in timeline:
        tc = f"{timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}"
        board.append(f"| {tc} | {scene['duration']}초 | {scene['id']} | {scene.get('slide', '별도 영상')} | {cell(scene.get('title', scene.get('name', '')))} |")
    script = ["# 설명형 영상 촬영 대본", "",
              "> 이 문서에서 해야 할 일: 두 출연자가 발화와 자료 큐를 리허설하고, 장면별 실측을 editorial/rehearsal.csv에 기록한다. 대괄호와 촬영 지시는 읽지 않는다.", "", heading, "",
              "출연: " + " · ".join(name(role, film) for role in film["presenters"] if role != "전체"), "",
              "회사 인트로·AI 팀 소개·아웃트로에는 스튜디오 발화가 없다. 본편은 00:20에 시작한다. 자료 전환·화자 교대·시선·호흡·홀드는 장면 배정 시간에 포함한다.", ""]
    prompt = [f"{film['title']} — {film['subtitle']}", f"{film['version']} / {film['date']}",
              "발음 표기 대본 / 대괄호는 읽지 않는 자료 큐", ""]
    board += ["", "## 장면별 편집 지시", ""]
    for scene, row in zip(film["scenes"], rows, strict=True):
        tc = f"{timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}"
        speaker = " → ".join(name(role, film) for role in scene["speaker"].split(" → "))
        board += [f"### {scene['id']} · {scene['title']}", "",
                  f"- 자료: {scene['slide']} / {int(scene['slide'][1:])}쪽.",
                  f"- 등장 발화: **‘{scene['cue']['anchor']}’**.",
                  f"- 화면·효과: {scene['cue']['effect']}", f"- 촬영: {scene['camera']}",
                  f"- 원본: {', '.join(scene['assets']) or '키워드·도식'}.",
                  f"- 편집·홀드: {scene['hold']}", ""]
        script += [f"## SCENE #{int(scene['id'][1:])} · {scene['id']} / {scene['slide']}", "",
                   f"- **장면:** {scene['chapter']} · {tc} · {scene['duration']}초",
                   f"- **발화자:** {speaker}", f"- **자료:** PPTX {int(scene['slide'][1:])}쪽 ({scene['slide']})",
                   f"- **발화 추정:** {row['spoken_syllables']}음절 · 기준 {row['estimates_seconds']['baseline']:.1f}초 · 느림 {row['estimates_seconds']['slow']:.1f}초. 호흡 포함.", "",
                   "**스크립트**", ""]
        prompt += [f"[{scene['id']} / {scene['slide']} / {tc} / {scene['duration']}초]"]
        for turn in scene["turns"]:
            script += [f"**{name(turn['speaker'], film)}:** {cue_text(turn, scene, film)}", ""]
            prompt += [f"<{name(turn['speaker'], film)}>", cue_text(turn, scene, film, phonetic=True), ""]
        script += [f"- **장면 효과:** {scene['cue']['effect']}", f"- **촬영 지시:** {scene['camera']}",
                   f"- **참고 원본:** {', '.join(scene['assets']) or '키워드·도식'}",
                   f"- **편집·확인:** {scene['hold']}", ""]
        for cue in scene.get("secondary_cues", []):
            instruction = f"‘{cue['anchor']}’ → [{cue['label']}] {cue['effect']}"
            board += [f"- 추가 자료 큐: {instruction}", ""]
            script += [f"- **추가 자료 큐:** {instruction}", ""]
    shared = ["## 제작과 확인", "", film.get("source_basis", ""), "",
              "현재 내부 행사 사진·화면 원본은 미수급이다. 원본이 없으면 JSON에 정의된 키워드·도식과 스튜디오 화면으로 제작한다. 범용 외부 사진을 실제 사내 시설·행사 증거로 사용하지 않는다. 원본 수집과 사실 확인은 docs/03-facts-and-assets.md를 따른다.", "",
              "PPTX에는 실제 애니메이션이 없다. 자료 전환과 모션은 영상팀 편집 지시이며 각 장면 시간 안에서 구현한다. 숨김 부록 A01~A03은 본편 사용 제외지만 검토 PDF·PNG에는 포함된다.", "",
              "프로그램 사이의 연결은 교육문화의 관점을 설명한다. Job Fair가 Tech X를 발생시켰거나 세미나가 IAP 개발의 직접 원인이었다는 인과를 표현하지 않는다. 10월 세미나 상태는 개최 확인 후 갱신한다.", "",
              "content/editorial-film.json을 수정한 뒤 이 버전 빌더로 다시 생성한다. 기존 기본안과 월드컵안은 별도 원본·산출물이다.", ""]
    board += shared
    rates = " / ".join(f"{key} {value}" for key, value in film["reading_rates"].items())
    script += ["## 발화량과 리허설", "",
               f"총 {report['total_spoken_syllables']}음절. 본편 추정: 느림 {report['main_estimates_seconds']['slow']:.1f}초 / 기준 {report['main_estimates_seconds']['baseline']:.1f}초 / 빠름 {report['main_estimates_seconds']['fast']:.1f}초. 고정 영상 포함 느린 추정 {report['whole_film_estimates_seconds']['slow']:.1f}초.", "",
               f"속도 가정은 초당 음절 수 {rates}다. 영문·숫자를 발음으로 치환하며 화자명·대괄호 큐는 제외한다. 장면별 pause_seconds는 호흡 예산이다. 이 값은 제작 계획 추정이며 실제 리허설 속도나 실측을 대체하지 않는다.", "",
               "리허설 CSV의 measured_seconds에 장면별 실측을 기록한다. 대사·자료 큐·화자·발음·배정 시간이 같을 때만 이전 실측과 메모를 보존한다. 내용 해시가 바뀌면 해당 실측 칸을 비워 다시 측정한다. IAP는 대표 사례로 35~40초를 유지하고, 초과 시 다른 장면의 중복 설명부터 조정한다.", "",
               "10분 촬영 발화 상한은 사용 가능한 한 테이크와 고정 영상 기준이다. NG·준비·추가 테이크의 총 촬영 시간은 별도다.", ""] + shared
    (DOCS / "01-storyboard.md").write_text("\n".join(board).rstrip() + "\n", encoding="utf-8")
    (DOCS / "02-script.md").write_text("\n".join(script).rstrip() + "\n", encoding="utf-8")
    (OUT / "teleprompter.txt").write_text("\n".join(prompt).rstrip() + "\n", encoding="utf-8")
    rehearsal = OUT / "rehearsal.csv"
    saved = {}
    if rehearsal.is_file():
        with rehearsal.open(encoding="utf-8-sig", newline="") as handle:
            saved = {row["scene"]: row for row in csv.DictReader(handle)}
    with rehearsal.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["scene", "PPT", "start", "end", "allocated_seconds", "speaker", "syllables", "estimated_baseline_seconds", "measured_seconds", "notes", "content_hash"])
        for scene, row in zip(film["scenes"], rows, strict=True):
            fingerprint = scene_content_hash(scene, film)
            old = saved.get(scene["id"], {})
            if old.get("content_hash") != fingerprint:
                old = {}
            writer.writerow([scene["id"], scene["slide"], timecode(scene["start"]),
                             timecode(scene["start"] + scene["duration"]), scene["duration"], scene["speaker"],
                             row["spoken_syllables"], row["estimates_seconds"]["baseline"],
                             old.get("measured_seconds", ""), old.get("notes", ""), fingerprint])
    return report


def check_source(film):
    assert film["total_seconds"] == 300
    assert [s["id"] for s in film["scenes"]] == [f"S{i:02d}" for i in range(1, 11)]
    assert [s["slide"] for s in film["scenes"]] == [f"P{i:02d}" for i in range(1, 11)]
    assert [s["id"] for s in film["appendix"]] == [f"A{i:02d}" for i in range(1, 4)]
    assert sum(s["duration"] for s in film["scenes"]) == 275
    assert sorted((s["start"], s["duration"]) for s in film["fixed_segments"]) == [(0, 5), (5, 15), (295, 5)]
    assert set(film["reading_rates"]) == {"slow", "baseline", "fast"}
    assert 0 < film["reading_rates"]["slow"] <= film["reading_rates"]["baseline"] <= film["reading_rates"]["fast"]
    previous = 20
    for scene in film["scenes"]:
        assert scene["start"] == previous and scene["duration"] > 0, scene["id"]
        assert scene["pause_seconds"] >= 0, scene["id"]
        previous += scene["duration"]
        raw = " ".join(turn["text"] for turn in scene["turns"])
        for turn in scene["turns"]:
            assert turn["speaker"] in film["presenters"], scene["id"]
        for cue in [scene["cue"], *scene.get("secondary_cues", [])]:
            assert cue["anchor"] and raw.count(cue["anchor"]) == 1, scene["id"]
        assert measure(scene, film)["estimates_seconds"]["slow"] <= scene["duration"], scene["id"]
    assert previous == 295


def main():
    film = json.loads(SOURCE.read_text(encoding="utf-8"))
    check_source(film)
    prs = make_deck(film)
    OUT.mkdir(parents=True, exist_ok=True)
    DOCS.mkdir(parents=True, exist_ok=True)
    report = make_documents(film)
    prs.save(DECK)
    print(json.dumps({"slides": len(prs.slides), "spoken_syllables": report["total_spoken_syllables"],
                      "main_estimates_seconds": report["main_estimates_seconds"],
                      "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
