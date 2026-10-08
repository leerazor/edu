"""Create the editable deck and synchronized production documents."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
DOCS = ROOT / "docs"
FONT = "맑은 고딕"
COLORS = {
    "bg": "101D2A", "panel": "1B2D3E", "text": "F6F8F9",
    "muted": "B4C6D4", "mint": "83E8C5", "lilac": "C1B3F4",
    "line": "385568", "amber": "F0CB7B",
}
STAGES = ["밖에서 배우다", "서로를 이해하다", "함께 바꾸다", "다시 나누다"]


def timecode(seconds: int) -> str:
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def spoken(text: str, film: dict) -> str:
    for word, reading in sorted(film["pronunciations"].items(), key=lambda pair: -len(pair[0])):
        text = text.replace(word, reading)
    return text


def measure(scene: dict, film: dict) -> dict:
    raw = " ".join(turn["text"] for turn in scene["turns"])
    read = spoken(raw, film)
    # Hangul blocks are counted after explicit pronunciation expansion. Punctuation
    # is handled by the per-scene pause allowance; these are planning estimates.
    count = len(re.findall(r"[가-힣]", read))
    assert not re.search(r"[A-Za-z0-9]", read), f"Unmapped pronunciation: {scene['id']}"
    times = {key: round(count / rate + scene["pause_seconds"], 2)
             for key, rate in film["reading_rates"].items()}
    return {
        "id": scene["id"], "slide": scene["slide"], "raw_characters": len(raw),
        "spoken_syllables": count, "pause_seconds": scene["pause_seconds"],
        "allocated_seconds": scene["duration"], "estimates_seconds": times,
        "baseline_visual_hold_seconds": round(scene["duration"] - times["baseline"], 2),
    }


def name(role: str, film: dict) -> str:
    return film["presenters"].get(role, role)


def cue_text(turn: dict, scene: dict, film: dict, phonetic: bool = False) -> str:
    value = turn["text"]
    cues = [(scene["cue"]["anchor"], f"{scene['slide']} 등장")]
    cues += [(cue["anchor"], cue["label"]) for cue in scene.get("secondary_cues", [])]
    for anchor, label in cues:
        if anchor in value:
            value = value.replace(anchor, f"[{label}] {anchor}", 1)
    if not phonetic:
        return value
    # Keep production cues unchanged; phonetic expansion applies only to speech.
    return "".join(piece if piece.startswith("[") else spoken(piece, film)
                   for piece in re.split(r"(\[[^\]]*\])", value))


def make_documents(film: dict) -> dict:
    measures = [measure(scene, film) for scene in film["scenes"]]
    report = {
        "basis": "Planning estimate, not a rehearsal measurement",
        "version": film["version"], "date": film["date"],
        "reading_rates_syllables_per_second": film["reading_rates"],
        "main_allocated_seconds": sum(scene["duration"] for scene in film["scenes"]),
        "fixed_seconds": sum(part["duration"] for part in film["fixed_segments"]),
        "scenes": measures,
        "total_spoken_syllables": sum(row["spoken_syllables"] for row in measures),
        "main_estimates_seconds": {
            key: round(sum(row["estimates_seconds"][key] for row in measures), 2)
            for key in film["reading_rates"]
        },
    }
    report["whole_film_estimates_seconds"] = {
        key: round(value + report["fixed_seconds"], 2)
        for key, value in report["main_estimates_seconds"].items()
    }
    (OUT / "timing-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")

    board = ["# 장면별 콘티와 편집 타임라인", "",
             "> 이 문서에서 해야 할 일: 최종 5분의 장면·자료·촬영·효과를 확인하고 내부 원본을 확보한 뒤, 편집에서 발화 큐에 맞춰 화면을 전환한다.", "",
             f"**검토용 초안 {film['version']} · {film['date']}**. 타임코드는 최종 편집 목표이며 실측 발화 시점이 아니다. 본편은 275초, 고정 영상은 25초다.", "",
             "## 전체 타임라인", "",
             "| 최종 타임코드 | 길이 | 장면 | PPTX 페이지 | 발화자 | 화면 목적 |",
             "| --- | ---: | --- | --- | --- | --- |",
             "| 00:00–00:05 | 5초 | F00 회사 인트로 | 없음 | 발화 없음 | 회사 제공 영상 |",
             "| 00:05–00:20 | 15초 | F01 AI 팀 소개 | 없음 | 발화 없음 | 별도 제작 브리프의 3컷 |"]
    script = ["# 촬영용 대본", "",
              "> 이 문서에서 해야 할 일: 류경동 팀장·이관우 교육담당이 장면별 발화를 리허설하고, 대괄호 자료 큐에 맞춰 촬영한다. 읽는 문장과 촬영 지시를 구분하고 실측 시간은 리허설 표에 기록한다.", "",
              f"**검토용 초안 {film['version']} · {film['date']}**. 제목: **AI 4 AI — 연결이 성장의 인프라가 되다**.", "",
              "회사 인트로·AI 소개·회사 아웃트로에는 스튜디오 발화가 없다. 본편은 00:20에 시작한다. 장면 시간에는 말 사이 호흡과 자료 홀드를 포함한다. 자료는 발화 큐를 기준으로 편집하고 초 단위 자동 자막으로 취급하지 않는다.", "",
              "**출연 자막:** 류경동 / Data Center팀 팀장 · 이관우 / Data Center팀 교육담당. 인사 직후 첫 개별 화면에서 각각 3초 표시한다.", "",
              "**읽기:** AI 4 AI = 에이아이 포 에이아이, IAP = 아이에이피, Tech X = 테크 엑스, Job Fair = 잡 페어, DCIS = 디씨아이에스. 공식 영문 확장명은 P02에 보여 주고 전부 읽지는 않는다.", ""]
    prompt = ["2026 Learning Fair | AI 4 AI", "연결이 성장의 인프라가 되다",
              "검토용 발음 표기 대본 / 대괄호는 읽지 않는 자료 큐",
              "출연: 류경동 팀장, 이관우 교육담당", ""]
    for scene, row in zip(film["scenes"], measures, strict=True):
        tc = f"{timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}"
        speaker = " → ".join(name(role, film) for role in scene["speaker"].split(" → "))
        board.append(f"| {tc} | {scene['duration']}초 | {scene['id']} {scene['chapter']} | {scene['slide']} / {int(scene['slide'][1:])}쪽 | {speaker} | {scene['title']} |")
        script += [f"## SCENE #{int(scene['id'][1:])} · {scene['id']} / {scene['slide']}", "",
                   f"- **장면:** {scene['chapter']} ({scene['duration']}초, 최종 {tc})",
                   f"- **발화자:** {speaker}",
                   f"- **자료:** PPTX {int(scene['slide'][1:])}쪽 ({scene['slide']})",
                   f"- **발화 추정:** {row['spoken_syllables']}음절, 기준 {row['estimates_seconds']['baseline']:.1f}초 / 느린 속도 {row['estimates_seconds']['slow']:.1f}초. 여유는 자료·시선·전환에 사용.", "",
                   "**스크립트**", ""]
        prompt += [f"[{scene['id']} / {scene['slide']} / 최종 {tc} / {scene['duration']}초]"]
        for turn in scene["turns"]:
            script += [f"**{name(turn['speaker'], film)}:** {cue_text(turn, scene, film)}", ""]
            prompt += [f"<{name(turn['speaker'], film)}>", cue_text(turn, scene, film, phonetic=True), ""]
        script += [f"- **장면 효과:** {scene['cue']['effect']}",
                   f"- **촬영 지시:** {scene['camera']}",
                   f"- **참고 원본:** {', '.join(scene['assets']) or '편집 가능한 키워드·도식'}",
                   f"- **편집·확인:** {scene['hold']}", ""]
        for cue in scene.get("secondary_cues", []):
            script += [f"- **추가 큐:** ‘{cue['anchor']}’에서 {cue['effect']}"]
        script += [""]
    board += ["| 04:55–05:00 | 5초 | F02 회사 아웃트로 | 없음 | 발화 없음 | 회사 제공 영상 |",
              "| **00:00–05:00** | **300초** | **고정 25초 + 본편 275초** | **본편 10쪽** | | |", "",
              "## 장면별 편집 지시", "",
              "본편은 활동의 시간순 기록이 아닌 교육문화의 연결을 보여 준다. 성장 회로는 문화 모델이다. 특히 Job Fair가 Tech X를 발생시켰거나, 외부 세미나가 IAP의 직접 원인이었다는 인과를 표현하지 않는다.", ""]
    for scene in film["scenes"]:
        board += [f"### {scene['id']} · {scene['title']}", "",
                  f"- 자료 등장 발화: **‘{scene['cue']['anchor']}’** → {scene['slide']}.",
                  f"- 화면·효과: {scene['cue']['effect']}",
                  f"- 카메라: {scene['camera']}",
                  f"- 원본: {', '.join(scene['assets']) or '별도 사진 불필요'}.",
                  f"- 지시: {scene['hold']}", ""]
        for cue in scene.get("secondary_cues", []):
            board += [f"- 추가 자료 큐 ‘{cue['anchor']}’: {cue['effect']}", ""]
    board += ["## 편집 규격과 다음 작업", "",
              "- 16:9 기준. 최종 해상도·프레임레이트는 영상팀과 맞춘다.",
              "- 연결선 점등·키워드 순차 등장은 영상팀 구현 지시다. 현재 PPTX에는 애니메이션을 넣지 않았다.",
              "- A01~A03은 참고 부록이며 본편 타임라인에 포함하지 않는다. PPTX에서도 숨김 슬라이드로 표시한다.",
              "- 각 장면별 최종 시간을 지키되 발화가 빠르면 실제 원본·타이틀 홀드로 채운다. 원본이 없을 때는 동일 메시지의 키워드 컷을 쓴다.",
              "- 컷 사이 1초 겹침·디졸브를 넣어 러닝타임이 늘어나면 해당 장면 안에서 길이를 조정한다.",
              "- 촬영 전 실제 10월 세미나 개최 여부와 원본 수집 목록을 확인하고, 리허설 실측으로 발화 추정을 교체한다.", ""]
    script += ["## 발화량과 리허설", "",
               f"총 발음 환산 **{report['total_spoken_syllables']}음절**. 호흡 포함 본편 추정은 느린 속도 {report['main_estimates_seconds']['slow']:.1f}초, 기준 {report['main_estimates_seconds']['baseline']:.1f}초, 빠른 속도 {report['main_estimates_seconds']['fast']:.1f}초다. 고정 영상 25초를 포함해도 10분 촬영 발화 상한 안에 들어간다.", "",
               "추정은 초당 4.2 / 4.6 / 5.0음절을 가정한 제작 계획이다. 실제 화자 속도를 측정한 결과가 아니다. 숫자·영문을 발음으로 치환하고 대괄호 큐·화자명은 발화량에서 제외했다. 발화자 교대·강조·마침표 쉼은 장면별 최소 호흡 예산으로 묶었다.", "",
               "**리허설 절차:** 장면별로 한 번 읽고 `deliverables/rehearsal.csv`의 실측 칸에 초를 적는다. IAP 핵심 장면은 먼저 유지하고, 초과하면 다른 사례의 보조 문장을 먼저 압축한다. 2회 촬영한다면 사용 가능한 한 테이크를 기준으로 발화량을 판단한다. NG·준비 시간은 최종 영상 길이에 합산하지 않는다.", "",
               "**읽지 않는 항목:** 장면 제목, 타임코드, 대괄호 큐, 설명·효과·카메라·원본 지시. 회사 제공 인트로/아웃트로와 AI 소개 영상은 별도 삽입한다.", "",
               "**수정 방법:** `content/film.json`의 발화·큐·시간을 수정한 뒤 빌드한다. 이 문서와 콘티·PPT 발표자 노트·프롬프터·시간 보고서는 자동 갱신된다. 콘셉트·사실·전달 문서는 영향이 있을 때 함께 수정한다.", ""]
    (DOCS / "05-storyboard.md").write_text("\n".join(board).rstrip() + "\n")
    (DOCS / "06-script.md").write_text("\n".join(script).rstrip() + "\n")
    (OUT / "teleprompter.txt").write_text("\n".join(prompt).rstrip() + "\n")
    rehearsal_path = OUT / "rehearsal.csv"
    existing_rehearsal = {}
    if rehearsal_path.exists():
        with rehearsal_path.open(encoding="utf-8-sig", newline="") as handle:
            existing_rehearsal = {row["scene"]: row for row in csv.DictReader(handle)}
    with rehearsal_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["scene", "PPT", "start", "end", "allocated_seconds", "speaker", "syllables", "estimated_baseline_seconds", "measured_seconds", "notes"])
        for scene, row in zip(film["scenes"], measures, strict=True):
            saved = existing_rehearsal.get(scene["id"], {})
            writer.writerow([scene["id"], scene["slide"], timecode(scene["start"]), timecode(scene["start"] + scene["duration"]), scene["duration"], scene["speaker"], row["spoken_syllables"], row["estimates_seconds"]["baseline"], saved.get("measured_seconds", ""), saved.get("notes", "")])
    return report


def rgb(key: str) -> RGBColor:
    return RGBColor.from_string(COLORS.get(key, key))


def text(slide, value: str, x: float, y: float, w: float, h: float,
         size: int = 24, color: str = "text", bold: bool = False,
         align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        east_asia = OxmlElement("a:ea")
        east_asia.set("typeface", FONT)
        run._r.get_or_add_rPr().append(east_asia)
    return shape


def box(slide, x, y, w, h, fill="panel", line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def connector(slide, x1, y1, x2, y2, color="mint", width=2, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        tip = OxmlElement("a:tailEnd")
        tip.set("type", "triangle")
        shape.line._get_or_add_ln().append(tip)
    return shape


def placeholder(slide, x, y, w, h, asset, caption):
    box(slide, x, y, w, h, line="line", radius=True)
    text(slide, "실제 원본 삽입", x + .25, y + .16, w - .5, .28, 13, "mint", True)
    text(slide, caption, x + .25, y + (.61 if h < 2.3 else .94), w - .5, .74,
         19 if h < 2.3 else 23, bold=True)
    text(slide, asset, x + .25, y + h - .35, w - .5, .26, 11, "muted")


def base(prs, page, label, title, active=None, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("bg")
    box(slide, .6, .53, .08, .3, "mint")
    text(slide, "2026 LEARNING FAIR  /  DATA CENTER", .86, .5, 9.7, .35, 12, "muted")
    text(slide, page, 11.64, .49, .85, .35, 13, "mint", True, PP_ALIGN.RIGHT)
    text(slide, label, .75, 1.13, 11.85, .4, 16, "mint", True)
    text(slide, title, .75, 1.67, 11.85, 1.25, 32, bold=True)
    if subtitle:
        text(slide, subtitle, .75, 2.78, 11.85, .45, 17, "muted")
    connector(slide, .75, 6.6, 12.58, 6.6, "line", 1)
    if active is not None:
        for i, stage in enumerate(STAGES):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.8 + i * 3.02), Inches(6.93), Inches(.11), Inches(.11))
            dot.fill.solid()
            dot.fill.fore_color.rgb = rgb("mint" if i <= active else "line")
            dot.line.fill.background()
            text(slide, stage, .99 + i * 3.02, 6.85, 2.65, .35, 12,
                 "text" if i == active else "muted", i == active)
    else:
        text(slide, "검토용 초안 · 회사 KV/폰트 수급 후 교체", .75, 6.86, 11.9, .35, 12, "muted")
    return slide


def note(slide, scene, film):
    lines = [f"{scene['slide']} / {scene['id']} / 최종 {timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}",
             "자료 큐는 발화 기준. 타임코드는 목표 편집 시간.",
             f"자료 등장: {scene['cue']['anchor']}", scene["cue"]["effect"],
             f"촬영: {scene['camera']}", f"원본: {', '.join(scene['assets']) or '키워드·도식'}",
             "스크립트:"]
    for turn in scene["turns"]:
        lines.append(f"{name(turn['speaker'], film)}: {cue_text(turn, scene, film)}")
    lines += [scene["hold"], "현재 애니메이션 없음. 연결선 점등 등은 영상팀 편집 지시.",
              "출처: 사용자 브리프 및 추가 확인, 2026-10-09. 외부 기술 배경은 docs/04-web-references.md."]
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def make_deck(film):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = f"{film['title']} — {film['subtitle']}"
    prs.core_properties.subject = "2026 Learning Fair 촬영·리디자인용 초안"
    prs.core_properties.author = "Data Center팀 교육담당 이관우"
    prs.core_properties.comments = "본편 10장, 참고 부록 3장. 내부 사진 자리표시자. 회사 KV 미적용."

    s = base(prs, "P01", "AI 시대, 우리의 질문", "AI를 움직이는 인프라,\n인프라를 움직이는 사람.")
    for x, value, caption in [(0.75, "약 160명", "Data Center팀"), (4.77, "서버·클라우드", "구축 · 제공 · 운영"), (8.79, "AI 인프라", "회사 AI 활용의 기반")]:
        box(s, x, 3.55, 3.77, 1.77, radius=True)
        text(s, value, x + .24, 3.88, 3.3, .65, 27, "mint", True)
        text(s, caption, x + .24, 4.64, 3.3, .4, 17, "muted")
    text(s, "우리 자신은 어떻게 성장해야 할까요?", .75, 5.65, 11.8, .55, 24, bold=True)

    s = base(prs, "P02", "AI 4 AI", "AI를 위한 인프라,\nAI로 발전하는 운영.")
    box(s, .75, 3.32, 5.37, 1.66, radius=True)
    box(s, 7.21, 3.32, 5.37, 1.66, radius=True)
    text(s, "AI를 위한 인프라 제공", 1.02, 3.86, 4.85, .75, 25, "mint", True)
    text(s, "AI를 활용한 운영", 7.55, 3.86, 4.7, .75, 25, "lilac", True)
    connector(s, 6.25, 3.84, 7.04, 3.84, arrow=True)
    connector(s, 7.04, 4.37, 6.25, 4.37, "lilac", arrow=True)
    text(s, "그 사이를 잇는 힘은, 함께 배우는 사람", .75, 5.27, 11.8, .55, 23, bold=True)
    text(s, film["official_expansion"], .75, 6.02, 11.8, .35, 14, "muted")

    s = base(prs, "P03", "밖에서 배우다", "밖의 지식이\n우리의 질문을 넓힙니다.", 0)
    text(s, "월 1회", .75, 3.51, 4.5, .8, 39, "mint", True)
    text(s, "외부 전문가와 만나는 기술 세미나", .75, 4.38, 5.1, .65, 20)
    text(s, "AI · 인프라 · 새로운 컴퓨팅", .75, 5.17, 5.1, .45, 18, "muted")
    text(s, "약 2천 명의 센터원 대상", .75, 5.79, 5.1, .4, 18, "muted")
    placeholder(s, 6.19, 3.29, 6.39, 2.93, "IMG02 · IMG03", "세미나 교류 사진 / 포스터")

    s = base(prs, "P04", "서로를 이해하다", "부서 사이의 거리를 줄였습니다.", 1)
    for x, title, body, asset, image in [(.75, "Job Fair · 11개 부서", "서로의 일을 직접 소개", "IMG04 · IMG05", "부스 전경 / A0 판넬"), (6.8, "DCIS · 월례 공유", "업무 · 서비스 · 연구 기술", "IMG06", "대표 세션 사진 / 발표 표지")]:
        text(s, title, x, 3.04, 5.6, .45, 23, "mint", True)
        text(s, body, x, 3.61, 5.6, .35, 18, "muted")
        placeholder(s, x, 4.11, 5.78, 2.07, asset, image)

    s = base(prs, "P05", "함께 바꾸다 / Tech X", "이해에서,\n함께 해보는 경험으로.", 2)
    text(s, "6개", .75, 3.36, 2.5, .8, 44, "mint", True)
    text(s, "프로젝트", .75, 4.22, 2.5, .4, 20, "muted")
    text(s, "44명", 3.3, 3.36, 2.6, .8, 44, "lilac", True)
    text(s, "참여 중", 3.3, 4.22, 2.6, .4, 20, "muted")
    box(s, 6.18, 3.38, 6.4, 1.39, radius=True)
    text(s, "원래 부서 밖의 관심 프로젝트에 참여", 6.49, 3.81, 5.8, .64, 22, bold=True)
    text(s, "새로운 경험  +  부서 간 협업", .75, 5.37, 11.8, .55, 24, bold=True)
    text(s, "IMG07: 모집·선발 공지 / 참여 집계 원본 확인", .75, 6.07, 11.8, .3, 12, "muted")

    s = base(prs, "P06", "함께 바꾸다 / IAP", "사람이 하던 알림 업무를 자동화", 2,
             "Tech X 과제 → 팀 주요 미션 → 운영 도입")
    text(s, "IAP  /  Infra Agent Platform", .75, 3.42, 11.8, .5, 23, "mint", True)
    for i, (title, caption) in enumerate([("로그 패턴 감지", "서버 로그의 특정 패턴"), ("정형 메시지", "정해진 포맷으로 구성"), ("메신저 동보", "알림 자동 전달")]):
        x = .75 + i * 4.08
        box(s, x, 4.17, 3.69, 1.37, radius=True)
        text(s, title, x + .23, 4.43, 3.23, .46, 24, "mint" if i != 1 else "lilac", True)
        text(s, caption, x + .23, 5.01, 3.23, .32, 15, "muted")
        if i < 2:
            connector(s, x + 3.75, 4.84, x + 4.02, 4.84, arrow=True)
    text(s, "현재 사례: 수작업 대체   /   확대 방향: 팀 인프라 전반", .75, 5.88, 11.8, .42, 19, bold=True)
    # The real interface is requested in notes, rather than invented on-screen.

    s = base(prs, "P07", "다시 나누다 / IAP", "만든 사람이,\n함께 쓰는 사람에게.", 3)
    for x, title, body, asset, image in [(.75, "6월 · 팀 소개", "개발 경험을 공유", "IMG09", "IAP 소개 세미나"), (6.8, "7월 · 실무 실습", "담당자 25명과 활용 방법 익히기", "IMG10", "Agent 등록 실습")]:
        text(s, title, x, 3.2, 5.6, .5, 26, "mint", True)
        text(s, body, x, 3.81, 5.6, .45, 18, "muted")
        placeholder(s, x, 4.48, 5.78, 1.69, asset, image)

    s = base(prs, "P08", "다시 나누다 / 명장 투어", "경험은 현장에서도 전해집니다.", 3)
    text(s, "7월 · 8명", .75, 3.17, 5.15, .75, 38, "mint", True)
    text(s, "명장과 함께한 화성 HPC센터", .75, 4.16, 5.15, .77, 22, bold=True)
    text(s, "시험 중 수냉식 서버\n공냉식 서버랙", .75, 5.21, 5.15, .9, 19, "muted")
    placeholder(s, 6.19, 3.23, 6.39, 2.95, "IMG11 · IMG12", "실제 투어 / 서버 체험 사진")

    s = base(prs, "P09", "다시 나누다 / 지식 자산과 확산 제안", "한 번의 배움을,\n다음 사람의 출발점으로.", 3)
    text(s, "AI 4 AI 녹화 → 교육 플랫폼 공유", .75, 3.23, 11.8, .5, 22, "mint", True)
    for i, word in enumerate(["배움", "이해", "실행", "공유"]):
        x = .75 + i * 3.08
        box(s, x, 4.0, 2.59, .95, radius=True)
        text(s, word, x + .15, 4.21, 2.29, .5, 25, "mint", True, PP_ALIGN.CENTER)
        if i < 3:
            connector(s, x + 2.7, 4.48, x + 2.96, 4.48, arrow=True)
    connector(s, 11.29, 5.07, 11.29, 5.3, "lilac")
    connector(s, 11.29, 5.3, 2.04, 5.3, "lilac")
    connector(s, 2.04, 5.3, 2.04, 5.07, "lilac", arrow=True)
    text(s, "다른 조직의 시작 제안", .75, 5.58, 11.8, .35, 15, "muted")
    text(s, "교류 열기 → 작은 공동 과제 → 다시 공유", .75, 6.02, 11.8, .43, 22, bold=True)

    s = base(prs, "P10", "우리의 성장 방식", "AI 4 AI")
    # Keep the closing card spare; the central title is deliberately larger.
    title_box = s.shapes[4]
    title_box.text_frame.paragraphs[0].runs[0].font.size = Pt(58)
    text(s, film["subtitle"], .75, 3.57, 11.8, .9, 35, "mint", True)
    text(s, "함께 배우고 · 함께 실행하고 · 다시 나눕니다", .75, 5.21, 11.8, .55, 23, "muted")
    text(s, "Data Center팀", .75, 5.99, 11.8, .42, 19)

    for slide, scene in zip(prs.slides, film["scenes"], strict=True):
        note(slide, scene, film)

    s = base(prs, "A01", "참고 부록 / 본편 사용 제외", "AI 4 AI 세미나 참고 목록")
    seminars = [
        ("3월", "Cisco", "AIOps"), ("4월", "NVIDIA", "GTC 2026 · 신규 제품 · 전략"),
        ("5월", "KISTI", "국가슈퍼컴퓨터 6호기 · 슈퍼컴 트렌드"),
        ("6월", "Pure Storage", "AI Ready Data 스토리지 플랫폼"),
        ("7월", "윤덕환 박사", "AI 시대 외로운 개인 · 동반자 AI"),
        ("8월", "류훈 금오공대 교수", "양자컴퓨터"),
        ("9월", "xFusion", "중국 AI 산업"),
        ("10월 · 예정", "OutSystems", "AI 시대의 Vibe Code Platform 전략"),
    ]
    for i, (month, host, topic) in enumerate(seminars):
        y = 3.02 + i * .39
        if i % 2 == 0:
            box(s, .75, y, 11.83, .38)
        text(s, month, .92, y + .03, 1.47, .31, 14, "amber" if i == 7 else "mint", True)
        text(s, host, 2.52, y + .03, 3.0, .31, 14)
        text(s, topic, 5.65, y + .03, 6.58, .31, 14)
    text(s, "10월은 2026-10-09 현재 미개최 · 촬영 전 개최 예정", .75, 6.22, 11.8, .3, 12, "amber")
    s.notes_slide.notes_text_frame.text = "참고 부록 A01. 본편에 삽입하지 않음. 출처: 사용자 브리프 및 추가 확인. 10월 개최 확인 후 상태와 날짜 수정. 월별 포스터 IMG03과 대조. 외부 공식 참고 docs/04-web-references.md."

    s = base(prs, "A02", "제작 체크 / 본편 사용 제외", "원본 자료 수집과 최종 확인")
    items = [
        ("IAP 실질 성과", "IMG08 · 실제 로그 감지 / 메신저 알림 화면", "실제 적용 범위 확인 · 계정/IP/서버명 등 가림"),
        ("교류 장면", "IMG02~07 · 세미나 / 부스 / DCIS / Tech X", "대표 사진과 공식 명칭 · 숫자 최신성 확인"),
        ("실습·현장·기록", "IMG09~13 · IAP 실습 / HPC 투어 / 플랫폼", "실제 행사 원본 · 시설/인물 공개 범위 확인"),
        ("최종 제작", "IMG14 · 회사 인트로/아웃트로 / KV / 폰트", "10월 개최 상태 갱신 · 리허설 실측 · 본편 275초"),
    ]
    for i, (title, body, detail) in enumerate(items):
        y = 3.01 + i * .78
        text(s, title, .75, y, 2.75, .38, 18, "mint", True)
        text(s, body, 3.54, y, 8.84, .37, 17)
        text(s, detail, 3.54, y + .39, 8.84, .3, 13, "muted")
    s.notes_slide.notes_text_frame.text = "참고 부록 A02. 본편에 삽입하지 않음. 상세 원본 목록 docs/03-facts-and-assets.md. 내부 원본 현재 미수급. 실제 효과 수치·후기 미제공. 전체 인프라 확대는 목표. 출연: 류경동 팀장, 이관우 교육담당."

    s = base(prs, "A03", "외부 시각 참고 / 본편 사용 제외", "범용 서버랙·냉각 참고 이미지")
    manifest = json.loads((ROOT / "assets/references/manifest.json").read_text())
    for i, entry in enumerate(manifest[:2]):
        x = .75 + i * 6.05
        picture = s.shapes.add_picture(str(ROOT / entry["file"]), Inches(x), Inches(3.17), width=Inches(5.76), height=Inches(2.61))
        image_ratio = picture.image.size[0] / picture.image.size[1]
        target_ratio = 5.76 / 2.61
        if image_ratio > target_ratio:
            picture.crop_left = picture.crop_right = (1 - target_ratio / image_ratio) / 2
        else:
            picture.crop_top = picture.crop_bottom = (1 - image_ratio / target_ratio) / 2
        caption = "NERSC 서버랙, 2011" if i == 0 else "1U 서버 공냉 구성, 2011"
        text(s, caption, x, 5.9, 5.76, .36, 17, bold=True)
        text(s, f"{entry['author']} · {entry['license']}", x, 6.32, 5.76, .27, 11, "muted")
    s.notes_slide.notes_text_frame.text = "\n".join([
        "참고 부록 A03. 실제 팀/화성 HPC센터/명장 투어 사진이 아님. 본편에 삽입하지 않음.",
        "CC0 이미지 2장. 이 장표에서는 레이아웃에 맞춰 크기를 조정함. 원본 JPEG는 저장소에 변경 없이 보관.",
        *[f"{entry['author']} / {entry['license']} / {entry['source_page']} / {entry['license_url']}" for entry in manifest[:2]],
        "전체 외부 출처 docs/04-web-references.md, assets/references/manifest.json.",
    ])
    # Explicitly hidden in slide show; they remain visible/editable in PowerPoint.
    for slide in list(prs.slides)[10:]:
        slide._element.set("show", "0")
    prs.save(OUT / "learning-fair-2026.pptx")
    return len(prs.slides)


def main():
    OUT.mkdir(exist_ok=True)
    film = json.loads((ROOT / "content/film.json").read_text())
    assert sum(scene["duration"] for scene in film["scenes"]) == 275
    assert sum(part["duration"] for part in film["fixed_segments"]) == 25
    previous = 20
    for scene in film["scenes"]:
        assert scene["start"] == previous
        previous += scene["duration"]
        full = " ".join(turn["text"] for turn in scene["turns"])
        assert full.count(scene["cue"]["anchor"]) == 1
        for cue in scene.get("secondary_cues", []):
            assert full.count(cue["anchor"]) == 1
    assert previous == 295
    report = make_documents(film)
    count = make_deck(film)
    print(json.dumps({"slides": count, "spoken_syllables": report["total_spoken_syllables"],
                      "main_estimates_seconds": report["main_estimates_seconds"],
                      "PPTX": "deliverables/learning-fair-2026.pptx"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
