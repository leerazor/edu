"""Build the additional tournament cut without touching the original package."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_package import box, connector, cue_text, measure, name, rgb, text, timecode

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/worldcup"
DOCS = ROOT / "docs/worldcup"


def make_documents(film):
    rows = [measure(scene, film) for scene in film["scenes"]]
    report = {
        "basis": "Planning estimate, not a rehearsal measurement",
        "version": film["version"], "date": film["date"],
        "reading_rates_syllables_per_second": film["reading_rates"],
        "main_allocated_seconds": sum(s["duration"] for s in film["scenes"]),
        "fixed_seconds": sum(s["duration"] for s in film["fixed_segments"]),
        "scenes": rows,
        "total_spoken_syllables": sum(r["spoken_syllables"] for r in rows),
        "main_estimates_seconds": {
            key: round(sum(r["estimates_seconds"][key] for r in rows), 2)
            for key in film["reading_rates"]
        },
    }
    report["whole_film_estimates_seconds"] = {
        key: round(value + report["fixed_seconds"], 2)
        for key, value in report["main_estimates_seconds"].items()
    }
    report["iap_allocated_seconds"] = film["scenes"][5]["duration"]
    (OUT / "timing-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    disclaimer = film["tournament"]["disclaimer"]
    board = ["# 월드컵 버전 콘티", "",
             "> 이 문서에서 해야 할 일: 4강·결승의 선택과 자료 큐를 확인하고, 각 장면 안에서 진출 효과·대화·자료 홀드를 편집한다.", "",
             f"**{film['version']} · {film['date']} · 기존안과 별도 제작본.** {disclaimer}", "",
             "전체 300초 = 회사 인트로 5초 + AI 팀 소개 15초 + 스튜디오 275초 + 회사 아웃트로 5초. P01~P10/S01~S10은 이 버전 폴더 안에서만 대응한다.", "",
             f"**공통 선택 기준:** {film['tournament']['criterion']}", "",
             "| 최종 타임코드 | 길이 | 장면 | 자료 | 화면 목적 |",
             "| --- | ---: | --- | --- | --- |",
             "| 00:00–00:05 | 5초 | F00 회사 인트로 | 회사 제공 | 고정 |",
             "| 00:05–00:20 | 15초 | F01 AI 팀 소개 | 별도 영상 | 기존 AI 소개 브리프 사용 |"]
    script = ["# 월드컵 버전 촬영 대본", "",
              "> 이 문서에서 해야 할 일: 두 출연자가 선택 이유를 소리 내어 읽고 자신의 의견에 맞는지 확인한 뒤, 대괄호 큐와 함께 리허설한다. 실측은 이 버전의 rehearsal.csv에 기록한다.", "",
              f"**{film['version']} · {film['date']}** · {film['title']} — {film['subtitle']}", "",
              f"**진행 규칙:** {film['tournament']['criterion']} 4강 두 경기와 결승 한 경기를 진행한다.", "",
              f"**선택의 성격:** {disclaimer}", "",
              "**출연 자막:** 류경동 / Data Center팀 팀장 · 이관우 / Data Center팀 교육담당. 첫 개별 화면에서 각각 3초 표시한다.", "",
              "팀장님은 선택의 기준과 업무 활용, 배움의 재공유를 묻고 자기 판단을 설명한다. 교육담당은 실제 운영 사실을 답하고 자신의 선택 이유도 말한다. 실제 행사에 직접 참석했거나 프로그램을 지시했다는 회고는 추가하지 않는다.", "",
              "회사 인트로·15초 AI 소개·아웃트로에는 스튜디오 발화가 없다. 대괄호·효과·촬영 지시는 읽지 않는다. 타임코드는 최종 편집 목표이며 실측 발화 시점이 아니다.", ""]
    prompter = [f"{film['title']} — {film['subtitle']}", "월드컵 추가 버전 / 대괄호는 읽지 않는 큐",
                "류경동 팀장 · 이관우 교육담당", f"제작 참고(읽지 않음): {disclaimer}", ""]
    for scene, row in zip(film["scenes"], rows, strict=True):
        tc = f"{timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}"
        board.append(f"| {tc} | {scene['duration']}초 | {scene['id']} {scene['chapter']} | {scene['slide']} / {int(scene['slide'][1:])}쪽 | {scene['title']} |")
        script += [f"## SCENE #{int(scene['id'][1:])} · {scene['id']} / {scene['slide']}", "",
                   f"- **장면:** {scene['chapter']} · {tc} · {scene['duration']}초",
                   f"- **발화자:** {scene['speaker']}",
                   f"- **자료:** 월드컵 PPTX {int(scene['slide'][1:])}쪽 ({scene['slide']})",
                   f"- **발화 추정:** {row['spoken_syllables']}음절 · 기준 {row['estimates_seconds']['baseline']:.1f}초 · 느린 속도 {row['estimates_seconds']['slow']:.1f}초. 호흡 포함.", "",
                   "**스크립트**", ""]
        prompter += [f"[{scene['id']} / {scene['slide']} / {tc} / {scene['duration']}초]"]
        for turn in scene["turns"]:
            script += [f"**{name(turn['speaker'], film)}:** {cue_text(turn, scene, film)}", ""]
            prompter += [f"<{name(turn['speaker'], film)}>", cue_text(turn, scene, film, phonetic=True), ""]
        script += [f"- **장면 효과:** {scene['cue']['effect']}",
                   f"- **촬영 지시:** {scene['camera']}",
                   f"- **참고 원본:** {', '.join(scene['assets']) or '편집 가능한 키워드·도식'}",
                   f"- **편집·확인:** {scene['hold']}", ""]
        for cue in scene.get("secondary_cues", []):
            script += [f"- **추가 자료 큐:** ‘{cue['anchor']}’ → [{cue['label']}] {cue['effect']}", ""]
    board += ["| 04:55–05:00 | 5초 | F02 회사 아웃트로 | 회사 제공 | 고정 |",
              "| **00:00–05:00** | **300초** | **고정 25초 + 본편 275초** | **본편 10쪽** | |", "",
              "## 장면별 촬영·모션", ""]
    for scene in film["scenes"]:
        board += [f"### {scene['id']} · {scene['title']}", "",
                  f"- **‘{scene['cue']['anchor']}’** → {scene['slide']} 등장.",
                  f"- 효과: {scene['cue']['effect']}", f"- 촬영: {scene['camera']}",
                  f"- 원본: {', '.join(scene['assets']) or '키워드·도식'}",
                  f"- 편집: {scene['hold']}", ""]
        for cue in scene.get("secondary_cues", []):
            board += [f"- **‘{cue['anchor']}’** → [{cue['label']}] {cue['effect']}", ""]
    shared = ["## 제작과 수정", "",
              "현재 PPTX는 편집 가능한 정적 자료다. 진출선 강조·카드 이동·결승 공개는 영상팀 구현 큐이며 PPTX에 실제 애니메이션은 없다. 자료를 읽을 때는 화면을 멈추고, 효과음은 대사를 덮지 않게 한다. 전환 시간은 각 장면에 포함한다.", "",
              "**결승 선택 공개:** 본편 P08은 선택 후 완성 상태다. ‘결승입니다’에서는 deliverables/worldcup/choice-preview/slide-01.png를 먼저 보여 주고, ‘저는 Tech X입니다’에서 slide-02.png 또는 본편 P08로 전환한다. 두 상태의 편집 원본은 choice-states.pptx다. 추가 장면이 아니라 같은 S08의 24초 안에서 교체하는 자료다.", "",
              "IAP 소개·실제 로그 알림·소개 세미나·실습은 S06의 40초 안에 묶었다. 다른 장면으로 IAP를 더 늘리지 않는다. 결승은 Tech X의 부서 간 경험과 참여 방식을 기준으로 결정한다.", "",
              "내부 원본은 미수급이다. 원본이 없으면 동일 메시지의 키워드 화면과 스튜디오 컷을 쓴다. A01~A03은 숨김 제작 부록이며 본편에 넣지 않는다. 외부 범용 사진을 내부 행사 증거로 대체하지 않는다.", "",
              "선택 결과를 바꾸려면 content/worldcup-film.json의 tournament와 해당 발화를 함께 수정하고, scripts/build_worldcup_package.py에서 선택 이유 장표까지 검토해 재생성한다. 기존 content/film.json 및 기존 제작 파일은 별도 버전이다.", "",
              "촬영 전 10월 OutSystems 개최 여부와 참여 수치의 최신성, 내부 사진, 실제 출연자 선택 이유를 확인한다. 10월은 마지막 확인 기준 미개최·촬영 전 개최 예정이다.", ""]
    board += shared
    script += ["## 발화량과 리허설", "",
               f"총 **{report['total_spoken_syllables']}음절**. 본편 추정: 느림 {report['main_estimates_seconds']['slow']:.1f}초 / 기준 {report['main_estimates_seconds']['baseline']:.1f}초 / 빠름 {report['main_estimates_seconds']['fast']:.1f}초. 고정 영상 25초 포함 느린 추정 {report['whole_film_estimates_seconds']['slow']:.1f}초. 남는 시간은 반응·자료·선택 공개에 사용한다.", "",
               "초당 4.2/4.6/5.0음절을 가정하고 영문·숫자를 발음으로 바꿔 계산했다. 장면별 호흡은 포함하지만 실제 발화 속도는 아직 측정하지 않았다. 리허설 CSV에 장면별 실측과 메모를 기록하고, 초과하면 중복 설명·긴 반응부터 줄인다. 3경기의 선택 이유와 교류 설명을 남긴다.", "",
               "10분 상한은 이 대본의 사용 가능한 한 테이크와 고정 영상 기준으로 확인한다. NG·준비·추가 테이크의 총 촬영 시간은 별도다.", ""] + shared
    (DOCS / "01-storyboard.md").write_text("\n".join(board).rstrip() + "\n")
    (DOCS / "02-script.md").write_text("\n".join(script).rstrip() + "\n")
    (OUT / "teleprompter.txt").write_text("\n".join(prompter).rstrip() + "\n")
    rehearsal = OUT / "rehearsal.csv"
    saved = {}
    if rehearsal.exists():
        with rehearsal.open(encoding="utf-8-sig", newline="") as handle:
            saved = {row["scene"]: row for row in csv.DictReader(handle)}
    with rehearsal.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["scene", "PPT", "start", "end", "allocated_seconds", "speaker", "syllables", "estimated_baseline_seconds", "measured_seconds", "notes"])
        for scene, row in zip(film["scenes"], rows, strict=True):
            old = saved.get(scene["id"], {})
            writer.writerow([scene["id"], scene["slide"], timecode(scene["start"]), timecode(scene["start"] + scene["duration"]), scene["duration"], scene["speaker"], row["spoken_syllables"], row["estimates_seconds"]["baseline"], old.get("measured_seconds", ""), old.get("notes", "")])
    return report


def base(prs, page, label, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("bg")
    box(slide, .65, .5, .1, .29, "mint")
    text(slide, "AI 4 AI  /  우리 팀 배움 월드컵", .9, .47, 10.9, .4, 14, "muted")
    text(slide, page, 11.8, .47, .85, .4, 14, "mint", True, PP_ALIGN.RIGHT)
    text(slide, label, .7, 1.16, 11.9, .4, 17, "mint", True)
    text(slide, title, .7, 1.78, 11.9, .88, 33, bold=True)
    if subtitle:
        text(slide, subtitle, .7, 2.72, 11.9, .55, 17, "muted")
    connector(slide, .7, 6.78, 12.63, 6.78, "line", 1)
    text(slide, "2026 LEARNING FAIR  /  DATA CENTER", .7, 7.03, 7.5, .25, 10, "muted")
    text(slide, "검토용 구성안 · 선택 결과는 촬영용 제안", 7.0, 7.0, 5.6, .32, 11, "muted", align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, detail=None, accent="mint"):
    box(slide, x, y, w, h, radius=True)
    box(slide, x, y, .055, h, accent)
    text(slide, heading, x + .27, y + .23, w - .54, .52, 24, accent, True)
    text(slide, body, x + .27, y + .98, w - .54, 1.2, 22, bold=True)
    if detail:
        text(slide, detail, x + .27, y + h - .7, w - .54, .53, 15, "muted")


def bracket(slide, film, reveal=False, winner_revealed=False):
    candidates = {c["id"]: c["name"] for c in film["tournament"]["candidates"]}
    matches = film["tournament"]["matches"]
    ys = [3.42, 4.12, 5.15, 5.85]
    for index, candidate in enumerate(matches[0]["candidates"] + matches[1]["candidates"]):
        box(slide, .75, ys[index], 3.37, .55, radius=True)
        text(slide, candidates[candidate], .93, ys[index] + .11, 3.0, .35, 17, bold=True)
    for index, match in enumerate(matches[:2]):
        first = ys[index * 2] + .275
        second = ys[index * 2 + 1] + .275
        center = (first + second) / 2
        connector(slide, 4.12, first, 4.55, first, "line")
        connector(slide, 4.12, second, 4.55, second, "line")
        connector(slide, 4.55, first, 4.55, second, "line")
        connector(slide, 4.55, center, 5.07, center, "mint", arrow=True)
        box(slide, 5.12, center - .28, 3.1, .56, radius=True)
        value = candidates[match["winner"]] if reveal else f"준결승 {index + 1}의 선택"
        finalist = text(slide, value, 5.27, center - .15, 2.8, .35, 16, "mint" if reveal else "muted", True)
        finalist.name = "bracket-final-left" if index == 0 else "bracket-final-right"
        connector(slide, 8.22, center, 8.72, center, "line")
    connector(slide, 8.72, 4.045, 8.72, 5.775, "line")
    connector(slide, 8.72, 4.91, 9.1, 4.91, "mint", arrow=True)
    box(slide, 9.16, 4.4, 3.4, 1.02, radius=True)
    value = candidates[matches[2]["winner"]] if winner_revealed else "오늘의 추천"
    winner = text(slide, value, 9.33, 4.68, 3.06, .45, 23, "mint", True, PP_ALIGN.CENTER)
    winner.name = "bracket-winner"
    if winner_revealed:
        reason = text(slide, "새로운 일에 함께 참여", 9.16, 5.65, 3.4, .42, 16, "muted", align=PP_ALIGN.CENTER)
        reason.name = "bracket-winner-reason"
    text(slide, "4강", .75, 3.02, 3.1, .3, 12, "muted")
    text(slide, "결승", 5.12, 3.02, 3.1, .3, 12, "muted")


def final_slide(prs, film, revealed):
    slide = base(prs, "P08", "결승 / 서로 알아보기 vs 함께 해보기", "Job Fair vs Tech X", film["tournament"]["criterion"])
    bracket(slide, film, reveal=True, winner_revealed=revealed)
    return slide


def make_deck(film):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = f"{film['title']} — {film['subtitle']}"
    prs.core_properties.subject = "월드컵 추가 버전 · 촬영용 선택안 · 300초"
    prs.core_properties.author = "Data Center팀 교육담당 이관우"
    prs.core_properties.comments = "본편 10장, 숨김 부록 3장. 실제 투표 결과 아님. 내부 원본·회사 KV 미수급."
    candidates = {c["id"]: c["name"] for c in film["tournament"]["candidates"]}
    matches = film["tournament"]["matches"]
    criterion = film["tournament"]["criterion"]

    s = base(prs, "P01", "오늘의 질문", "AI 4 AI · 우리 팀 배움 월드컵", criterion)
    text(s, "AI를 위한 인프라를 만들고,\nAI로 운영을 발전시키는 팀.", .75, 3.62, 7.2, 1.4, 28, bold=True)
    box(s, 9.0, 3.58, 3.56, 1.83, radius=True)
    text(s, "약 160명", 9.26, 3.84, 3.04, .65, 34, "mint", True)
    text(s, "서버 · 클라우드 구축과 운영", 9.26, 4.71, 3.04, .52, 16, "muted")
    text(s, "우리의 성장 방식은, 서로 배우는 것에서 시작합니다.", .75, 5.85, 11.7, .5, 23, "mint", True)
    text(s, film["official_expansion"], .75, 6.4, 11.7, .26, 12, "muted")

    s = base(prs, "P02", "대진표 / 4가지 교류 경험", "세 번의 선택, 하나의 추천", criterion)
    bracket(s, film)

    s = base(prs, "P03", "준결승 1 / 외부 전문가와 동료", "새로운 관점 vs 서로의 업무 이해", criterion)
    card(s, .75, 3.58, 5.72, 2.7, "AI 4 AI 세미나", "외부 전문가와 월 1회\nAI·인프라를 함께 학습", "센터원 대상 · 녹화 후 교육 플랫폼 공유")
    card(s, 6.84, 3.58, 5.72, 2.7, "Job Fair", "11개 부서의 일을\n직접 묻고 알아가는 시간", "8월 · 부스에서 A0 판넬로 업무 소개", "lilac")
    text(s, film["official_expansion"], .75, 6.4, 11.81, .26, 12, "muted")

    winner = candidates[matches[0]["winner"]]
    s = base(prs, "P04", "준결승 1 / 선택 이유", f"결승 진출 · {winner}", "다른 조직도 동료의 일을 소개하는 자리부터")
    card(s, .75, 3.56, 5.72, 2.63, "Job Fair", "부서의 일을 직접 듣고\n궁금한 점을 질문", "실제 자료: 11개 부서 부스 · 업무 판넬")
    card(s, 6.84, 3.56, 5.72, 2.63, "DCIS", "주요 업무·서비스·연구를\n월 1회 나눔", "정기적인 내부 교류도 함께 운영", "lilac")
    text(s, "같은 교류의 목적을 가진 두 활동", .75, 6.33, 11.7, .3, 14, "muted")

    s = base(prs, "P05", "준결승 2 / 현장에서 배우고, 함께 해보고", "명장의 경험 vs 새로운 프로젝트", criterion)
    card(s, .75, 3.58, 5.72, 2.78, "명장 투어", "7월 · 8명 대상\n화성 HPC센터 현장 학습", "시험 중 수냉식 서버 · 공냉식 서버랙")
    card(s, 6.84, 3.58, 5.72, 2.78, "Tech X", "6개 프로젝트 · 44명\n다른 부서의 과제에 참여", "원래 부서에서 하기 어려운 새로운 경험", "lilac")

    s = base(prs, "P06", "Tech X / 대표 성과 사례 한 가지", "함께 만든 것을, 함께 쓰는 경험으로", "IAP · Infra Agent Platform / 팀 주요 미션으로 발전 · 운영 도입")
    for i, (title, description) in enumerate([("로그 패턴 감지", "서버 로그의 특정 패턴"), ("정형 메시지", "정해진 포맷으로 구성"), ("메신저 동보", "사람이 하던 알림 업무 대체")]):
        x = .75 + i * 4.08
        box(s, x, 3.6, 3.65, 1.38, radius=True)
        text(s, title, x + .23, 3.88, 3.19, .43, 23, "mint", True)
        text(s, description, x + .23, 4.48, 3.19, .34, 14, "muted")
        if i < 2:
            connector(s, x + 3.74, 4.29, x + 4.0, 4.29, arrow=True)
    box(s, .75, 5.24, 11.81, .72, radius=True)
    text(s, "6월 · 팀 소개 세미나     /     7월 · 실무 담당자 25명 실습", 1.0, 5.43, 11.3, .38, 21, "lilac", True)
    text(s, "현재: 로그 알림 수작업 대체   |   확대 목표: 팀 인프라 전반", .75, 6.24, 11.81, .34, 16, "muted")

    winner = candidates[matches[1]["winner"]]
    s = base(prs, "P07", "준결승 2 / 선택 이유", f"결승 진출 · {winner}", "평소 맡지 못한 과제를, 다른 부서 동료와 직접 해볼 기회")
    for i, (heading, body) in enumerate([("관심", "해보고 싶은 과제"), ("참여", "부서 밖의 동료"), ("경험", "함께 실행하는 배움")]):
        x = .75 + i * 4.08
        card(s, x, 3.58, 3.65, 2.31, heading, body, accent="mint" if i != 1 else "lilac")
    text(s, "프로젝트의 성과와, 구성원이 배울 기회를 함께 봅니다.", .75, 6.25, 11.81, .4, 22, bold=True)

    s = final_slide(prs, film, revealed=True)

    s = base(prs, "P09", "다른 조직에 전하는 시작 제안", "작은 공동 과제부터 시작해 보세요", "Tech X의 참여 방식을 다른 조직에 맞게 적용하는 제안")
    for i, (heading, body) in enumerate([("01  과제 열기", "도움이 필요한\n작은 과제"), ("02  함께 하기", "관심 있는\n다른 부서 동료"), ("03  다시 나누기", "해본 경험과\n배운 점 공유")]):
        x = .75 + i * 4.08
        card(s, x, 3.58, 3.65, 2.7, heading, body, accent="mint" if i != 1 else "lilac")
    text(s, "다른 조직의 도입 실적이 아닌, 시작 방법의 제안입니다.", .75, 6.42, 11.81, .25, 12, "muted")

    s = base(prs, "P10", "AI를 위한 인프라, 함께 배우는 사람", "AI 4 AI — 배움이 부서를 넘을 때")
    text(s, "서로의 지식과 경험이\n우리의 다음 성장을 만듭니다.", .75, 3.36, 11.8, 1.5, 34, "mint", True)
    for i, label in enumerate(["외부 전문가", "다른 부서 동료", "함께하는 과제", "현장의 경험"]):
        x = .75 + i * 3.02
        box(s, x, 5.38, 2.75, .76, radius=True)
        text(s, label, x + .1, 5.59, 2.55, .4, 18, align=PP_ALIGN.CENTER)
    text(s, film["official_expansion"], .75, 6.4, 11.8, .26, 12, "muted")

    for slide, scene in zip(prs.slides, film["scenes"], strict=True):
        lines = [f"{scene['slide']} / {scene['id']} / 최종 {timecode(scene['start'])}–{timecode(scene['start'] + scene['duration'])}",
                 film["tournament"]["disclaimer"], f"자료 등장: {scene['cue']['anchor']}", scene["cue"]["effect"],
                 f"촬영: {scene['camera']}", f"원본: {', '.join(scene['assets']) or '키워드·도식'}", "스크립트:"]
        lines += [f"{name(turn['speaker'], film)}: {cue_text(turn, scene, film)}" for turn in scene["turns"]]
        lines += [f"추가 큐: {cue['anchor']} / {cue['label']} / {cue['effect']}" for cue in scene.get("secondary_cues", [])]
        lines += [scene["hold"], "현재 PPTX 애니메이션 없음. 모션은 영상팀 편집 지시.",
                  "출처: 사용자 브리프·추가 확인. 사실 기준 docs/03-facts-and-assets.md. 원본 미수급 시 현재 키워드 화면 사용."]
        slide.notes_slide.notes_text_frame.text = "\n".join(lines)

    s = base(prs, "A01", "숨김 부록 / 본편 사용 제외", "AI 4 AI 세미나 참고 목록")
    entries = [("3월", "Cisco", "AIOps"), ("4월", "NVIDIA", "GTC 2026 · 신규 제품 · 전략"),
               ("5월", "KISTI", "국가슈퍼컴퓨터 6호기 · 슈퍼컴 트렌드"),
               ("6월", "Pure Storage", "AI Ready Data 스토리지 플랫폼"),
               ("7월", "윤덕환 박사", "AI 시대 외로운 개인 · 동반자 AI"),
               ("8월", "류훈 금오공대 교수", "양자컴퓨터"), ("9월", "xFusion", "중국 AI 산업"),
               ("10월 · 예정", "OutSystems", "AI 시대의 Vibe Code Platform 전략")]
    for i, (month, host, topic) in enumerate(entries):
        y = 3.1 + i * .39
        if i % 2 == 0:
            box(s, .75, y, 11.81, .38)
        text(s, month, .93, y + .03, 1.45, .31, 14, "amber" if i == 7 else "mint", True)
        text(s, host, 2.52, y + .03, 3.0, .31, 14)
        text(s, topic, 5.65, y + .03, 6.56, .31, 14)
    text(s, "10월: 2026-10-09 마지막 확인 시 미개최 · 촬영 전 개최 예정", .75, 6.4, 11.81, .28, 12, "amber")
    s.notes_slide.notes_text_frame.text = "A01 / 사용자 브리프·추가 확인. 10월 미개최 상태는 개최 확인 후에만 갱신. IMG03 포스터와 대조. 본편 제외."

    s = base(prs, "A02", "숨김 부록 / 영상팀 준비", "실제 사진은 이 장면에 넣어 주세요")
    entries = [("P03~04", "IMG02~06 · IMG13", "세미나 · Job Fair · DCIS · 교육 플랫폼"),
               ("P05", "IMG07 · IMG11~12", "Tech X 참여 근거 · 명장 투어 / 서버 체험"),
               ("P06", "IMG08~10", "실제 로그 알림 · 6월 소개 · 7월 실습"),
               ("전체", "IMG01 · IMG14", "팀 소개 · 회사 인트로/아웃트로 · KV/폰트")]
    for i, (page, asset, description) in enumerate(entries):
        y = 3.18 + i * .73
        text(s, page, .75, y, 1.67, .4, 19, "mint", True)
        text(s, asset, 2.48, y, 3.66, .4, 17, bold=True)
        text(s, description, 6.2, y, 6.35, .5, 16)
    text(s, "내부 원본 미수급 · 현재 키워드로 진행 가능 · 외부 사진으로 실제 활동을 대체하지 않음", .75, 6.35, 11.81, .33, 13, "amber")
    s.notes_slide.notes_text_frame.text = "A02 / 원본 수집 목록 docs/03-facts-and-assets.md. 실제 내부 사진은 아직 없음. 공개 가능한 자료를 받고 민감정보를 가린 뒤 편집. 원본이 없으면 본편 키워드 도식 사용. 원본 코드의 장표 번호는 이 월드컵 버전의 매핑을 우선한다."

    s = base(prs, "A03", "숨김 부록 / 선택과 모션", "촬영 전에, 두 사람이 선택 이유를 맞춥니다")
    for i, (heading, body) in enumerate([
        ("현재 선택안", "Job Fair · Tech X 진출 → Tech X 추천 / 실제 투표 결과 아님"),
        ("다른 선택이라면", "대사·진출자·결승 대진·마지막 추천을 함께 수정 후 재생성"),
        ("모션 지시", "진출선 0.4초 강조 · 선택 공개 1초 홀드 · 설명 중 정지"),
        ("시간 기준", "본편 275초 · IAP는 S06의 40초 · 리허설로 실제 시간 확인"),
    ]):
        y = 3.25 + i * .72
        text(s, heading, .75, y, 2.32, .44, 20, "mint", True)
        text(s, body, 3.17, y, 9.39, .52, 18)
    text(s, "PPTX는 정적 자료입니다. 선택 효과와 사진 전환은 영상팀 편집으로 구현합니다.", .75, 6.35, 11.81, .34, 14, "muted")
    s.notes_slide.notes_text_frame.text = "A03 / 촬영용 선택안. 실제 선호를 확인하지 않았으므로 리허설에서 두 출연자가 판단을 검토한다. 모션 0.4초 및 홀드 1초는 각 장면 시간에 포함. 탈락 후보는 원래 가치를 유지하며 비하·폭발·가짜 점수 없이 처리. 기존docs/13-selection-show-format.md는 형식 비교 배경."
    for slide in list(prs.slides)[10:]:
        slide._element.set("show", "0")
    prs.save(OUT / "learning-fair-2026-worldcup.pptx")
    # Two actual editable states let the editor reveal the result at the speech
    # cue without briefly showing the finished main slide first.
    states = Presentation()
    states.slide_width, states.slide_height = prs.slide_width, prs.slide_height
    states.core_properties.title = "월드컵 P08 결승 선택 전·후 교체 화면"
    states.core_properties.author = prs.core_properties.author
    for revealed in [False, True]:
        state = final_slide(states, film, revealed)
        state.notes_slide.notes_text_frame.text = (
            f"P08 {'선택 후' if revealed else '선택 전'} / S08 동일 24초 안에서 교체. "
            + ("‘저는 Tech X입니다’ 발화에서 이 화면으로 전환." if revealed else "‘결승입니다’ 발화에서 이 화면을 먼저 표시. 최종 우승 표시는 아직 없음.")
            + " 제작자가 구성한 촬영용 선택안이며 실제 투표 결과 아님. 새로운 장면이나 추가 시간이 아니다."
        )
    states.save(OUT / "choice-states.pptx")
    return len(prs.slides)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    DOCS.mkdir(parents=True, exist_ok=True)
    film = json.loads((ROOT / "content/worldcup-film.json").read_text())
    assert sum(scene["duration"] for scene in film["scenes"]) == 275
    assert sum(part["duration"] for part in film["fixed_segments"]) == 25
    previous = 20
    for scene in film["scenes"]:
        assert scene["start"] == previous
        previous += scene["duration"]
        speech = " ".join(turn["text"] for turn in scene["turns"])
        for cue in [scene["cue"], *scene.get("secondary_cues", [])]:
            assert speech.count(cue["anchor"]) == 1, scene["id"]
        assert measure(scene, film)["estimates_seconds"]["slow"] <= scene["duration"], scene["id"]
    assert previous == 295
    report = make_documents(film)
    slides = make_deck(film)
    print(json.dumps({"slides": slides, "spoken_syllables": report["total_spoken_syllables"],
                      "main_estimates_seconds": report["main_estimates_seconds"], "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
